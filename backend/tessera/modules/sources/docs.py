"""ING-05 module: ``src.docs`` reads one XLSX report into a document with cell locators.

The MVP slice reads one report: the Retail Controlling weekly sales flash
(``sources/reports/retail_controlling/weekly_flash/``). The ``report`` param selects the file,
and ``DEFAULT_REPORT`` is the flash of the demo week when a run does not name one. One
``SourceDocument`` per workbook and one ``Segment`` per non-empty row: the segment's ``locator``
is the row's cell range (``Stores!A3:U3``), so evidence and metrics point back at the sheet
instead of at a wall of text. Sheets are announced by a ``[Sheet: ...]`` marker with the sheet as
locator. Only XLSX is read here; another format is a later task, not a universal parser.

Two facts come from the report's own words, never from a guess: the reporting period
(``Fiscal week 42: Mon 12/10/2026 - Sun 18/10/2026``) and the issue date (``Version 1.0 -
issued 20/10/2026``), the first line in workbook order that states them. They become
``event_start``/``event_end`` and ``available_at``, and ``business_date`` is the period end.
A workbook whose cover does not state them is skipped and reported: an undated report may carry
numbers from after ``as_of``.

Availability is a date, not an instant. ``generated/README.md`` Â§4 uses a report when its period
end and issue date are ``â©½`` the simulated day, so it is available from the start of the issue
date in the HQ timezone. The precision is that of the report, and ``availability_basis`` says so.

The canonical text shows the stored cell values, not the displayed ones: number formats (``%``,
``kâ¬``, thousands) are not emulated, so ``meta.text`` states it and a reader does not mistake
``0.1747`` for the file's ``17.5%``.
"""

import hashlib
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from tessera.core.config import get_settings
from tessera.core.contracts import ModuleManifest, ModuleResult, Segment, SourceDocument

DEFAULT_REPORT = "retail_controlling/weekly_flash/Flash_W42_2026.xlsx"
REPORTS_DIR = "reports"
DOC_VERSION = "v1"
TEXT_VERSION = "report-xlsx-v2"
REPORT_TIMEZONE = ZoneInfo("Europe/Rome")  # Milan HQ, the office that issues the flash
REPORT_LANGUAGE = "en"
TEXT_BASIS = (
    "cell values as stored, not as displayed: a share may be 0.1747 here and 17.5% in Excel"
)

MANIFEST = ModuleManifest(
    id="src.docs",
    version="1",
    layer="sources",
    description="One XLSX report as a document with one segment and cell locator per row",
    requires=(),
    optional_inputs=(),
    produces=("SourceDocument", "Segment"),
    enabled=True,
    mode="live",
    params_schema={
        "type": "object",
        "properties": {
            "report": {"type": "string", "minLength": 1},
            "language": {"type": "string", "pattern": "^[a-z]{2,3}(-[A-Za-z0-9]{2,8})*$"},
        },
        "additionalProperties": False,
    },
)

_ISSUED = re.compile(r"issued\s+(\d{1,2})/(\d{1,2})/(\d{4})", re.IGNORECASE)
_PERIOD = re.compile(
    r"fiscal\s+week\s+(\d{1,2})\s*:\s*(?:\w+\.?,?\s+)?(\d{1,2})/(\d{1,2})/(\d{4})"
    r"\s*[-\u2013]\s*(?:\w+\.?,?\s+)?(\d{1,2})/(\d{1,2})/(\d{4})",
    re.IGNORECASE,
)
_VERSION = re.compile(r"version\s+(\d+(?:\.\d+)?)", re.IGNORECASE)


class UndatedReport(ValueError):
    """The workbook is readable but does not state when it covers and when it was issued."""


@dataclass(frozen=True)
class ParsedReport:
    relative: str
    content_hash: str
    sheets: tuple[str, ...]
    rows: tuple[tuple[str, str], ...]  # (locator, text), workbook order
    fiscal_week: int
    version: str | None
    issue_date: date
    period_start: date
    period_end: date
    issued_line: str

    @property
    def available_at(self) -> datetime:
        day = max(self.issue_date, self.period_end)
        return datetime.combine(day, time(0, 0), tzinfo=REPORT_TIMEZONE)


def report_path(sources_dir: Path, relative: str) -> Path:
    """Resolve the selected report inside ``sources/reports``; refuse anything else."""
    root = (sources_dir / REPORTS_DIR).resolve()
    path = (root / relative).resolve()
    if not path.is_relative_to(root):
        raise ValueError("the report must live under sources/reports")
    if path.suffix.lower() != ".xlsx":
        raise ValueError("only .xlsx reports are read here")
    if not path.is_file():
        raise ValueError("no such report")
    return path


def parse_report(path: Path, *, relative: str) -> ParsedReport:
    """Read the workbook once: cover facts, then one row per non-empty spreadsheet row."""
    data = path.read_bytes()
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    try:
        issue_date, period_start, period_end, fiscal_week, version, issued_line = _cover(workbook)
        rows: list[tuple[str, str]] = []
        for worksheet in workbook.worksheets:
            rows.append((worksheet.title, f"[Sheet: {worksheet.title}]"))
            rows.extend(_rows(worksheet))
        sheets = tuple(worksheet.title for worksheet in workbook.worksheets)
    finally:
        workbook.close()
    return ParsedReport(
        relative=relative,
        content_hash=hashlib.sha256(data).hexdigest(),
        sheets=sheets,
        rows=tuple(rows),
        fiscal_week=fiscal_week,
        version=version,
        issue_date=issue_date,
        period_start=period_start,
        period_end=period_end,
        issued_line=issued_line,
    )


def _cover(workbook: Any) -> tuple[date, date, date, int, str | None, str]:
    """Find the reporting period and the issue date in the cells the report itself writes."""
    issued: tuple[date, str] | None = None
    period: tuple[int, date, date] | None = None
    version: str | None = None
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                text = cell.value
                if not isinstance(text, str):
                    continue
                if issued is None and (match := _ISSUED.search(text)):
                    issued = (_day(match.group(1), match.group(2), match.group(3)), text.strip())
                if period is None and (match := _PERIOD.search(text)):
                    period = (
                        int(match.group(1)),
                        _day(match.group(2), match.group(3), match.group(4)),
                        _day(match.group(5), match.group(6), match.group(7)),
                    )
                if version is None and (match := _VERSION.search(text)):
                    version = match.group(1)
    if issued is None:
        raise UndatedReport("no 'issued <date>' line on any sheet")
    if period is None:
        raise UndatedReport("no 'Fiscal week <n>: <start> - <end>' line on any sheet")
    return issued[0], period[1], period[2], period[0], version, issued[1]


def _day(day: str, month: str, year: str) -> date:
    try:
        return date(int(year), int(month), int(day))
    except ValueError as error:
        raise UndatedReport(f"the cover date is not a date: {error}") from None


def _rows(worksheet: Any) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for row in worksheet.iter_rows():
        values = [(cell.column, cell.value) for cell in row]
        filled = [column for column, value in values if _filled(value)]
        if not filled:
            continue
        first, last = min(filled), max(filled)
        text = " | ".join(_text(value) for column, value in values if first <= column <= last)
        locator = f"{worksheet.title}!{get_column_letter(first)}{row[0].row}"
        if first != last:
            locator += f":{get_column_letter(last)}{row[0].row}"
        rows.append((locator, text))
    return rows


def _filled(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or value.strip() != "")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return " ".join(value.split())  # a cell's own line breaks would break one row = one line
    if isinstance(value, bool):  # before int: a bool is an int
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return repr(value)  # shortest round-trip: the stored number, never a rounded one
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def document_id(relative: str) -> str:
    return "doc-" + Path(relative).stem.lower().replace("_", "-")


def build_document(
    report: ParsedReport, language: str = REPORT_LANGUAGE
) -> tuple[SourceDocument, list[Segment]]:
    document_id_ = document_id(report.relative)
    uri = f"dataset://{REPORTS_DIR}/{report.relative}"
    content_hash = f"sha256:{report.content_hash}"
    common = {
        "version": DOC_VERSION,
        "available_at": report.available_at,
        "availability_basis": "report_cover_issue_date_not_before_period_end",
    }
    text = "\n".join(text for _, text in report.rows)
    document = SourceDocument(
        id=document_id_,
        **common,
        source_refs=({"uri": uri, "content_hash": content_hash},),
        source_type="report",
        connector=MANIFEST.id,
        uri=uri,
        content_hash=content_hash,
        origin=report.relative.split("/")[0] if "/" in report.relative else None,
        author_ref=None,
        language=language,
        business_date=report.period_end,
        event_start=datetime.combine(report.period_start, time(0, 0), tzinfo=REPORT_TIMEZONE),
        event_end=datetime.combine(report.period_end, time(23, 59, 59), tzinfo=REPORT_TIMEZONE),
        thread_key=None,
        parent_document_id=None,
        text=text,
        text_version=TEXT_VERSION,
        meta={
            "report": report.relative,
            "format": "xlsx",
            "sheets": list(report.sheets),
            "fiscal_week": report.fiscal_week,
            "version": report.version,
            "issue_date": report.issue_date.isoformat(),
            "issued_line": report.issued_line,
            "text": TEXT_BASIS,
        },
    )
    segments: list[Segment] = []
    offset = 0
    for index, (locator, row_text) in enumerate(report.rows):
        segments.append(
            Segment(
                id=f"{document_id_}-seg{index}",
                **common,
                source_refs=({"uri": uri, "content_hash": content_hash, "locator": locator},),
                document_id=document_id_,
                text_version=TEXT_VERSION,
                idx=index,
                text=row_text,
                char_start=offset,
                char_end=offset + len(row_text),
                locator=locator,
                t_start_s=None,
                t_end_s=None,
                is_quoted=False,
            )
        )
        offset += len(row_text) + 1  # the "\n" joining this row to the next
    for segment in segments:
        segment.validate_against(document)
    return document, segments


def run(ctx, *, sources_dir: Path | None = None) -> ModuleResult:
    root = sources_dir or get_settings().sources_dir
    relative = str(ctx.params.get("report") or DEFAULT_REPORT)
    language = str(ctx.params.get("language") or REPORT_LANGUAGE)
    uri = f"dataset://{REPORTS_DIR}/{relative}"
    counts: dict[str, int] = {"documents": 0, "segments": 0, "sheets": 0}
    errors: list[str] = []
    abstentions: list[str] = []
    try:
        parsed = parse_report(report_path(root, relative), relative=relative)
    except UndatedReport as error:
        abstentions.append(f"{uri}: not read ({error})")
        counts["reports_skipped"] = 1
    except (ValueError, OSError, BadZipFile, InvalidFileException) as error:
        errors.append(f"{uri}: {error}")
    except Exception as error:  # openpyxl raises KeyError/ParseError on a malformed zip
        # A source file must never fail the run: the orchestrator would discard what the
        # modules already published.
        errors.append(f"{uri}: unreadable workbook ({type(error).__name__}: {error})")
    else:
        if parsed.available_at > ctx.spec.as_of:
            abstentions.append(f"{uri}: not issued at as_of")
            counts["not_available"] = 1
        else:
            document, segments = build_document(parsed, language)
            ctx.repo.put(document, *segments)
            counts.update(documents=1, segments=len(segments), sheets=len(parsed.sheets))
    return ModuleResult(
        module_id=MANIFEST.id,
        module_version=MANIFEST.version,
        status="degraded" if errors else "succeeded",
        counts=counts,
        errors=tuple(errors),
        abstentions=tuple(abstentions),
    )
