"""ING-06 module: ``ing.metrics`` turns one approved, versioned report mapping into facts.

AI reads, code counts. The mapping (``config/report_metrics.yaml``) names the report family, the
sheets, the header row, the key column and every column it maps. Code does the rest: it matches a
row's store name against ``erp/master_data/stores.csv``, reads the cell value from the workbook the
``SourceDocument`` came from, and publishes one ``MetricFact`` per metric and scope.

A number is never invented. A cell that is blank, textual (``n.a.``, ``closed``) or not finite is
not a zero and produces no fact. A mapped column whose definition sentence is missing from the
report's own text still publishes its facts, but as ``not_comparable``: a number without its
definition cannot be aggregated (RFC 4.5). The same holds for a reported EUR figure whose FX basis
the report does not state, and for an LFL figure on a store the report flags as not LFL.

The mapping is data, not code: ``mapping_version`` travels on every fact, ``status: approved`` is
required, and ``scale`` declares the conversion (EUR k to EUR, a stored fraction to a percent)
instead of letting a reader guess it. Only the explicit slice is read; the other sheets and the
other 23 reports are a later task, not a universal parser.

Aggregates are published only for a perimeter the run actually covers: a region row is written
only when every store of that region is in the run's stores, so a run scoped to MIL01 never leaks
an EMEA total that includes stores it was not allowed to read.

A corrected re-issue of the same week is a second document with its own id, so both versions are
facts. Reconciling them (``supersedes``) belongs to ``int.dedup``, not here.
"""

import csv
import hashlib
import math
import re
from collections import Counter
from dataclasses import dataclass
from fnmatch import fnmatchcase
from pathlib import Path
from typing import Any, Literal
from zipfile import BadZipFile

import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from tesseract.core.config import get_settings
from tesseract.core.contracts import (
    ArtifactRef,
    MetricFact,
    ModuleManifest,
    ModuleResult,
    RunSpec,
    SourceDocument,
    SourceRef,
)

DEFAULT_MAPPING = "report_metrics.yaml"
FACT_VERSION = "v1"
REPORTS_DIR = "reports"
AVAILABILITY_BASIS = "same_as_source_document"
UNITS = ("units", "customers", "requests", "visitors", "money", "percent")
# Physical counts: the units Quantity refuses to keep negative (a percent may fall).
COUNT_UNITS = ("units", "customers", "requests", "visitors")
LFL_FLAG = "Y"

MANIFEST = ModuleManifest(
    id="ing.metrics",
    version="1",
    layer="ingestion",
    description=(
        "One approved, versioned report mapping -> MetricFact with definition, period, "
        "currency/VAT and comparability"
    ),
    requires=("src.docs",),
    optional_inputs=(),
    produces=("MetricFact",),
    enabled=True,
    mode="live",
    params_schema={
        "type": "object",
        "properties": {"mapping": {"type": "string", "minLength": 1}},
        "additionalProperties": False,
    },
)

_CURRENCY = re.compile(r"[A-Z]{3}")
MetricUnit = Literal["units", "customers", "requests", "visitors", "money", "percent"]
Scope = Literal["store", "region"]


class MappingError(ValueError):
    """The mapping is missing, unapproved, or does not describe a usable slice."""


@dataclass(frozen=True)
class Column:
    header: str
    metric: str
    unit: str
    definition_quote: str
    currency: str | None = None
    currency_from: str | None = None
    scale: float = 1.0
    vat_included: bool | None = None
    gross_or_net: str | None = None
    fx_basis_quote: str | None = None
    requires_lfl: bool = False


@dataclass(frozen=True)
class Sheet:
    name: str
    header_row: int
    first_row: int
    key_column: str
    scope: str
    lfl_column: str | None
    columns: tuple[Column, ...]


@dataclass(frozen=True)
class Mapping:
    path: Path
    sha256: str
    version: int
    mapping_version: str
    report_pattern: str
    calendar: str | None
    sheets: tuple[Sheet, ...]


@dataclass(frozen=True)
class StoreRow:
    store_id: str
    name: str
    country_code: str
    region: str


@dataclass(frozen=True)
class Master:
    by_name: dict[str, StoreRow]
    regions: dict[str, frozenset[str]]  # region name -> every store id in it


@dataclass(frozen=True)
class Row:
    number: int
    values: dict[str, Any]


@dataclass(frozen=True)
class Table:
    headers: dict[str, int]  # normalized header -> 1-based column index
    rows: tuple[Row, ...]


def header(value: object) -> str:
    """A header cell as one line, so ``Net Sales\\nWk (EUR k)`` is readable in the mapping."""
    return " ".join(str(value).split())


def mapping_path(config_dir: Path, relative: str) -> Path:
    """Resolve the mapping inside ``config``; refuse anything else."""
    root = config_dir.resolve()
    path = (root / relative).resolve()
    if not path.is_relative_to(root):
        raise MappingError("the mapping must live under config")
    if path.suffix.lower() not in (".yaml", ".yml"):
        raise MappingError("the mapping must be a YAML file")
    if not path.is_file():
        raise MappingError("no such mapping")
    return path


def _fields(raw: dict[str, Any], allowed: set[str], where: str) -> None:
    unknown = set(raw) - allowed
    if unknown:
        raise MappingError(f"{where}: unknown key(s) {', '.join(sorted(unknown))}")


def _text(raw: dict[str, Any], key: str, where: str) -> str:
    value = raw.get(key)
    if not isinstance(value, str) or not value.strip():
        raise MappingError(f"{where}: {key} must be a non-empty string")
    return value.strip()


def _flag(raw: dict[str, Any], key: str, where: str) -> bool | None:
    value = raw.get(key)
    if value is None:
        return None
    if not isinstance(value, bool):
        raise MappingError(f"{where}: {key} must be true or false")
    return value


def load_mapping(path: Path) -> Mapping:
    data = path.read_bytes()
    try:
        raw = yaml.safe_load(data)
    except yaml.YAMLError as error:
        raise MappingError(f"not readable YAML ({error})") from None
    if not isinstance(raw, dict):
        raise MappingError("the mapping must be a YAML mapping")
    _fields(
        raw,
        {"version", "mapping_version", "status", "report_pattern", "calendar", "sheets"},
        "mapping",
    )
    if raw.get("status") != "approved":
        raise MappingError("the mapping is not approved")
    version = raw.get("version")
    if not isinstance(version, int) or isinstance(version, bool) or version < 1:
        raise MappingError("version must be a positive integer")
    calendar = raw.get("calendar")
    if calendar is not None and (not isinstance(calendar, str) or not calendar.strip()):
        raise MappingError("calendar must be a non-empty string when present")
    raw_sheets = raw.get("sheets")
    if not isinstance(raw_sheets, dict) or not raw_sheets:
        raise MappingError("sheets must be a non-empty mapping")
    return Mapping(
        path=path,
        sha256=hashlib.sha256(data).hexdigest(),
        version=version,
        mapping_version=_text(raw, "mapping_version", "mapping"),
        report_pattern=_text(raw, "report_pattern", "mapping"),
        calendar=calendar.strip() if calendar else None,
        sheets=tuple(
            _sheet(name, sheet) for name, sheet in raw_sheets.items() if isinstance(name, str)
        ),
    )


def _sheet(name: str, raw: Any) -> Sheet:
    where = f"sheet {name}"
    if not isinstance(raw, dict):
        raise MappingError(f"{where}: must be a mapping")
    _fields(
        raw,
        {"header_row", "first_row", "key_column", "scope", "lfl_column", "columns"},
        where,
    )
    header_row = raw.get("header_row")
    first_row = raw.get("first_row")
    for key, value in (("header_row", header_row), ("first_row", first_row)):
        if not isinstance(value, int) or isinstance(value, bool) or value < 1:
            raise MappingError(f"{where}: {key} must be a positive integer")
    if first_row <= header_row:
        raise MappingError(f"{where}: first_row must be below header_row")
    scope = raw.get("scope")
    if scope not in ("store", "region"):
        raise MappingError(f"{where}: scope must be 'store' or 'region'")
    lfl = raw.get("lfl_column")
    if lfl is not None and (not isinstance(lfl, str) or not lfl.strip()):
        raise MappingError(f"{where}: lfl_column must be a non-empty string when present")
    raw_columns = raw.get("columns")
    if not isinstance(raw_columns, dict) or not raw_columns:
        raise MappingError(f"{where}: columns must be a non-empty mapping")
    columns = tuple(
        _column(header(key), column, f"{where} column {key}")
        for key, column in raw_columns.items()
        if isinstance(key, str)
    )
    if len({column.header for column in columns}) != len(columns):
        raise MappingError(f"{where}: duplicate column headers")
    if len({column.metric for column in columns}) != len(columns):
        raise MappingError(f"{where}: two columns map to the same metric")
    if any(column.requires_lfl for column in columns) and lfl is None:
        raise MappingError(f"{where}: requires_lfl needs an lfl_column")
    return Sheet(
        name=name,
        header_row=header_row,
        first_row=first_row,
        key_column=header(_text(raw, "key_column", where)),
        scope=scope,
        lfl_column=header(lfl) if lfl else None,
        columns=columns,
    )


def _column(name: str, raw: Any, where: str) -> Column:
    if not isinstance(raw, dict):
        raise MappingError(f"{where}: must be a mapping")
    _fields(
        raw,
        {
            "metric",
            "unit",
            "definition_quote",
            "currency",
            "currency_from",
            "scale",
            "vat_included",
            "gross_or_net",
            "fx_basis_quote",
            "requires_lfl",
        },
        where,
    )
    unit = raw.get("unit")
    if unit not in UNITS:
        raise MappingError(f"{where}: unit must be one of {', '.join(UNITS)}")
    currency = raw.get("currency")
    currency_from = raw.get("currency_from")
    if currency is not None:
        if not isinstance(currency, str) or _CURRENCY.fullmatch(currency) is None:
            raise MappingError(f"{where}: currency must be a three-letter uppercase code")
    if currency_from is not None:
        currency_from = _text(raw, "currency_from", where)
    if unit == "money" and (currency is None) == (currency_from is None):
        raise MappingError(f"{where}: money needs exactly one of currency or currency_from")
    if unit != "money" and (currency is not None or currency_from is not None):
        raise MappingError(f"{where}: only money columns carry a currency")
    scale = raw.get("scale", 1.0)
    if (
        not isinstance(scale, (int, float))
        or isinstance(scale, bool)
        or not math.isfinite(scale)
        or scale <= 0
    ):
        raise MappingError(f"{where}: scale must be a positive finite number")
    vat_included = _flag(raw, "vat_included", where)
    gross_or_net = raw.get("gross_or_net")
    if gross_or_net is not None and gross_or_net not in ("gross", "net"):
        raise MappingError(f"{where}: gross_or_net must be 'gross' or 'net'")
    if unit != "money" and (vat_included is not None or gross_or_net is not None):
        raise MappingError(f"{where}: VAT and gross/net belong to money columns")
    fx_basis_quote = raw.get("fx_basis_quote")
    if fx_basis_quote is not None:
        fx_basis_quote = _text(raw, "fx_basis_quote", where)
        if unit != "money":
            raise MappingError(f"{where}: only money columns carry an FX basis")
    requires_lfl = _flag(raw, "requires_lfl", where) or False
    return Column(
        header=name,
        metric=_text(raw, "metric", where),
        unit=unit,
        definition_quote=_text(raw, "definition_quote", where),
        currency=currency,
        currency_from=currency_from,
        scale=float(scale),
        vat_included=vat_included,
        gross_or_net=gross_or_net,
        fx_basis_quote=fx_basis_quote,
        requires_lfl=requires_lfl,
    )


def load_master(sources_dir: Path) -> Master:
    """Store names and regions from the ERP master data, the only owner of store identity."""
    path = sources_dir / "erp" / "master_data" / "stores.csv"
    with path.open(encoding="utf-8", newline="") as stream:
        rows = [
            StoreRow(
                store_id=row["store_id"],
                name=row["store_name"],
                country_code=row.get("country_code") or "",
                region=row.get("region") or "",
            )
            for row in csv.DictReader(stream)
        ]
    regions: dict[str, set[str]] = {}
    for row in rows:
        regions.setdefault(row.region, set()).add(row.store_id)
    return Master(
        by_name={row.name: row for row in rows},
        regions={name: frozenset(ids) for name, ids in regions.items()},
    )


def read_workbook(path: Path, mapping: Mapping) -> dict[str, Table | None]:
    """One open per document; the mapped sheets only, ``None`` when a sheet is absent."""
    workbook = load_workbook(path, data_only=True)
    try:
        tables: dict[str, Table | None] = {}
        for sheet in mapping.sheets:
            tables[sheet.name] = (
                _table(workbook[sheet.name], sheet) if sheet.name in workbook.sheetnames else None
            )
        return tables
    finally:
        workbook.close()


def _table(worksheet: Any, sheet: Sheet) -> Table:
    headers: dict[str, int] = {}
    for cell in worksheet[sheet.header_row]:
        if cell.value is None or not str(cell.value).strip():
            continue
        name = header(cell.value)
        if name in headers:
            raise MappingError(f"sheet {sheet.name}: duplicate header {name!r}")
        headers[name] = cell.column
    by_column = {column: name for name, column in headers.items()}
    rows: list[Row] = []
    for row in worksheet.iter_rows(min_row=sheet.first_row):
        values = {by_column[cell.column]: cell.value for cell in row if cell.column in by_column}
        if _present(values.get(sheet.key_column)):
            rows.append(Row(number=row[0].row, values=values))
    return Table(headers=headers, rows=tuple(rows))


def _present(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or value.strip() != "")


def number(value: object) -> float | None:
    """The stored number, or ``None``. Text is never a number here, not even ``"589265"``."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    value = float(value)
    return value if math.isfinite(value) else None


def sentence(text: str, quote: str) -> str | None:
    """The report's own line that carries ``quote``, so a definition is quoted, not written."""
    for line in text.splitlines():
        if quote in line:
            return line.strip()
    return None


def scope_of(
    value: object, sheet: Sheet, master: Master
) -> tuple[dict[str, str], frozenset[str]] | None:
    """The fact's scope and its member stores, or ``None`` when the report names no known scope."""
    if not isinstance(value, str):
        return None
    name = " ".join(value.split())
    if sheet.scope == "store":
        store = master.by_name.get(name)
        if store is None:
            return None
        return (
            {
                "level": "store",
                "store_id": store.store_id,
                "region": store.region,
                "country_code": store.country_code,
            },
            frozenset((store.store_id,)),
        )
    members = master.regions.get(name) or master.regions.get(name.upper())
    if members is None:
        return None
    return {"level": "region", "region": name.upper()}, members


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def fact_id(document_id: str, scope: dict[str, str], metric: str) -> str:
    where = scope["store_id"] if scope["level"] == "store" else scope["region"]
    return f"mfx-{document_id.removeprefix('doc-')}-{_slug(where)}-{_slug(metric)}"


def build_facts(
    document: SourceDocument,
    mapping: Mapping,
    sheet: Sheet,
    table: Table,
    master: Master,
    spec: RunSpec,
) -> tuple[list[MetricFact], list[str], Counter]:
    notes: list[str] = []
    counts: Counter = Counter()
    assert document.event_start is not None and document.event_end is not None
    period_start, period_end = document.event_start, document.event_end
    if sheet.key_column not in table.headers:
        notes.append(f"{document.id}: key column {sheet.key_column!r} is missing from {sheet.name}")
        counts["columns_missing"] += 1
        return [], notes, counts
    present: list[Column] = []
    for column in sheet.columns:
        if column.header in table.headers:
            present.append(column)
        else:
            notes.append(f"{document.id}: column {column.header!r} is missing from {sheet.name}")
            counts["columns_missing"] += 1
    definitions: dict[str, str | None] = {}
    fx_bases: dict[str, str | None] = {}
    for column in present:
        definitions[column.metric] = sentence(document.text, column.definition_quote)
        if definitions[column.metric] is None:
            notes.append(f"{document.id}: {column.metric} has no definition in this report")
            counts["definitions_missing"] += 1
        fx_bases[column.metric] = (
            sentence(document.text, column.fx_basis_quote) if column.fx_basis_quote else None
        )
        if column.fx_basis_quote and fx_bases[column.metric] is None:
            notes.append(f"{document.id}: {column.metric} has no FX basis in this report")
            counts["fx_basis_missing"] += 1
    facts: list[MetricFact] = []
    for row in table.rows:
        scope = scope_of(row.values.get(sheet.key_column), sheet, master)
        if scope is None:
            notes.append(
                f"{document.id} {sheet.name}!{row.number}: "
                f"{row.values.get(sheet.key_column)!r} is not in the master data"
            )
            counts["rows_unknown_scope"] += 1
            continue
        scope_names, members = scope
        if not members <= set(spec.stores):
            where = scope_names.get("store_id") or scope_names.get("region") or sheet.scope
            notes.append(
                f"{document.id} {sheet.name}!{row.number}: {where} is outside this run's stores"
            )
            counts["rows_outside_run"] += 1
            continue
        lfl = sheet.lfl_column is not None and row.values.get(sheet.lfl_column) == LFL_FLAG
        for column in present:
            locator = f"{sheet.name}!{get_column_letter(table.headers[column.header])}{row.number}"
            raw = row.values.get(column.header)
            if not _present(raw):
                counts["cells_missing"] += 1
                continue
            value = number(raw)
            if value is None:
                notes.append(f"{document.id} {locator}: {raw!r} is not a number")
                counts["cells_not_numeric"] += 1
                continue
            currency = column.currency
            if column.currency_from is not None:
                found = row.values.get(column.currency_from)
                currency = found.strip() if isinstance(found, str) else None
                if currency is None or _CURRENCY.fullmatch(currency) is None:
                    notes.append(f"{document.id} {locator}: no usable currency for this row")
                    counts["cells_without_currency"] += 1
                    continue
            if column.unit in COUNT_UNITS and value < 0:
                notes.append(f"{document.id} {locator}: a {column.unit} count cannot be negative")
                counts["cells_negative"] += 1
                continue
            definition = definitions[column.metric]
            basis = fx_bases[column.metric]
            comparable = definition is not None and (basis is not None or not column.fx_basis_quote)
            if column.requires_lfl and not lfl:
                comparable = False
                counts["lfl_not_comparable"] += 1
            facts.append(
                MetricFact(
                    id=fact_id(document.id, scope_names, column.metric),
                    version=FACT_VERSION,
                    available_at=document.available_at,
                    availability_basis=AVAILABILITY_BASIS,
                    source_refs=(
                        SourceRef(
                            uri=document.uri,
                            content_hash=document.content_hash,
                            locator=locator,
                        ),
                    ),
                    run_id=spec.run_id,
                    scenario_id=spec.scenario_id,
                    computed_at=spec.as_of,
                    input_refs=(
                        ArtifactRef(
                            kind="SourceDocument", id=document.id, version=document.version
                        ),
                    ),
                    producer=MANIFEST.id,
                    producer_version=MANIFEST.version,
                    document_id=document.id,
                    metric=column.metric,
                    definition=definition,
                    value=value * column.scale,
                    unit=column.unit,  # type: ignore[arg-type]
                    currency=currency,
                    vat_included=column.vat_included,
                    gross_or_net=column.gross_or_net,  # type: ignore[arg-type]
                    scope=scope_names,
                    period_start=period_start,
                    period_end=period_end,
                    calendar=mapping.calendar,
                    fx_basis=basis,
                    locator=locator,
                    mapping_version=mapping.mapping_version,
                    comparability_status="comparable" if comparable else "not_comparable",
                )
            )
    return facts, notes, counts


def run(
    ctx,
    *,
    sources_dir: Path | None = None,
    config_dir: Path | None = None,
) -> ModuleResult:
    settings = get_settings()
    sources = sources_dir or settings.sources_dir
    config = config_dir or settings.config_dir
    relative_mapping = str(ctx.params.get("mapping") or DEFAULT_MAPPING)
    counts: Counter = Counter()
    errors: list[str] = []
    abstentions: list[str] = []
    try:
        mapping = load_mapping(mapping_path(config, relative_mapping))
        master = load_master(sources)
    except (OSError, MappingError) as error:
        return ModuleResult(
            module_id=MANIFEST.id,
            module_version=MANIFEST.version,
            status="failed",
            counts={},
            errors=(f"{relative_mapping}: {error}",),
            abstentions=(),
        )
    for document in sorted(ctx.repo.list(SourceDocument), key=lambda item: item.id):
        relative = document.meta.get("report")
        if not isinstance(relative, str) or document.meta.get("format") != "xlsx":
            abstentions.append(f"{document.id}: not an XLSX report")
            counts["documents_unmapped"] += 1
            continue
        if not fnmatchcase(relative, mapping.report_pattern):
            abstentions.append(f"{document.id}: {relative} is outside the approved mapping")
            counts["documents_unmapped"] += 1
            continue
        if document.event_start is None or document.event_end is None:
            abstentions.append(f"{document.id}: the report states no period")
            counts["documents_without_period"] += 1
            continue
        path = sources / REPORTS_DIR / relative
        try:
            tables = read_workbook(path, mapping)
        except (OSError, BadZipFile, InvalidFileException, MappingError) as error:
            errors.append(f"{document.id}: unreadable workbook ({type(error).__name__}: {error})")
            continue
        except Exception as error:  # openpyxl raises KeyError/ParseError on a malformed zip
            errors.append(f"{document.id}: unreadable workbook ({type(error).__name__})")
            continue
        written = 0
        for sheet in mapping.sheets:
            table = tables[sheet.name]
            if table is None:
                abstentions.append(f"{document.id}: sheet {sheet.name!r} is missing")
                counts["sheets_missing"] += 1
                continue
            facts, notes, stats = build_facts(document, mapping, sheet, table, master, ctx.spec)
            counts.update(stats)
            abstentions.extend(notes)
            if facts:
                ctx.repo.put(*facts)
            written += len(facts)
        counts["documents"] += 1
        counts["facts"] += written
    return ModuleResult(
        module_id=MANIFEST.id,
        module_version=MANIFEST.version,
        status="degraded" if errors else "succeeded",
        counts=dict(counts),
        errors=tuple(errors),
        abstentions=tuple(abstentions),
    )
