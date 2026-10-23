"""ing.metrics: the real W42 flash mapped into MetricFacts, and synthetic edges."""

from datetime import UTC, datetime
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from tessera.core.contracts import MetricFact, ModuleSnapshot, RunSpec
from tessera.core.db import connect
from tessera.core.orchestrator import run as orchestrate
from tessera.core.registry import RegisteredModule, Registry
from tessera.core.run_repo import repo_session
from tessera.modules.ingestion import metrics
from tessera.modules.ingestion.metrics import MANIFEST
from tessera.modules.sources import docs

ROOT = Path(__file__).resolve().parents[4]
SOURCES = ROOT / "generated" / "sources"
CONFIG = ROOT / "config"
FLASHES = "retail_controlling/weekly_flash"
AS_OF = datetime(2026, 10, 22, 20, 0, tzinfo=UTC)
STORE_HEADER = [
    "Plant",
    "Store",
    "Region",
    "Country",
    "Ccy",
    "Net Sales LC",
    "Net Sales EUR",
    "Traffic",
    "LFL",
    "\u0394% LFL cFX",
]
COVER = [
    ["MONCLER | Retail Controlling"],
    ["WEEKLY SALES FLASH - WEEK 42"],
    ["Fiscal week 42: Mon 12/10/2026 - Sun 18/10/2026"],
    ["Version 1.0 - issued 20/10/2026"],
    [],
    ["Definitions"],
    [
        "Net sales = gross sales - discounts - returns (returns booked on the processing "
        "date), excluding VAT/sales taxes."
    ],
    ["LFL = stores trading at least 5 days in both CY week and LY week"],
    ["Reported EUR at average October 2026 rates (month of the week's Thursday)"],
]
MASTER = """store_id,store_name,country_code,region
MIL01,Milano Montenapoleone,IT,EMEA
GVA01,Gen\u00e8ve Rue du Rh\u00f4ne,CH,EMEA
LON01,London New Bond Street,GB,EMEA
NYC01,New York Madison Avenue,US,AMERICAS
"""


def all_stores():
    return tuple(sorted(store.store_id for store in metrics.load_master(SOURCES).by_name.values()))


def spec(stores, *, as_of=AS_OF, report=docs.DEFAULT_REPORT, mapping=None, run_id="metrics-run"):
    params = {"mapping": mapping} if mapping else {}
    return RunSpec(
        run_id=run_id,
        scenario_id="historical",
        as_of=as_of,
        stores=tuple(stores),
        source_manifest_hash="sources",
        config_hash="config",
        code_version="test",
        mode="live",
        modules=(
            ModuleSnapshot(manifest=docs.MANIFEST, params={"report": report}),
            ModuleSnapshot(manifest=MANIFEST, params=params),
        ),
        prompt_versions={},
        model_versions={},
        authorized_fixture_refs=(),
    )


def execute(tmp_path, request, *, sources_dir=SOURCES, config_dir=CONFIG):
    engine = connect(f"sqlite:///{tmp_path / 'metrics.db'}")
    registry = Registry(
        (
            RegisteredModule(docs.MANIFEST, lambda ctx: docs.run(ctx, sources_dir=sources_dir)),
            RegisteredModule(
                MANIFEST,
                lambda ctx: metrics.run(ctx, sources_dir=sources_dir, config_dir=config_dir),
            ),
        )
    )
    statuses = orchestrate(request, engine=engine, registry=registry, clients={})
    return engine, {status.module_id: status for status in statuses}


def facts_of(tmp_path, request, *, sources_dir=SOURCES, config_dir=CONFIG):
    engine, statuses = execute(tmp_path, request, sources_dir=sources_dir, config_dir=config_dir)
    with repo_session(engine, request) as repo:
        return statuses[MANIFEST.id].result, repo.list(MetricFact)


def by_key(facts):
    where = lambda fact: fact.scope.get("store_id") or fact.scope["region"]  # noqa: E731
    return {(where(fact), fact.metric): fact for fact in facts}


def write_report(root, relative, sheets):
    path = root / "reports" / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row_index, values in enumerate(rows, start=1):
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row_index, column=column, value=value)
    book.save(path)
    return path


def stores_sheet(*rows):
    return [["Store detail"], STORE_HEADER, *[list(row) for row in rows]]


def summary_sheet(*rows):
    return [
        ["Weekly flash - Summary"],
        ["12/10 - 18/10/2026 | v1.0"],
        [],
        ["Region", "Net Sales Wk (EUR k)"],
        *[list(row) for row in rows],
    ]


def write_sources(tmp_path, *, sheets, relative=docs.DEFAULT_REPORT):
    root = tmp_path / "sources"
    (root / "erp" / "master_data").mkdir(parents=True, exist_ok=True)
    (root / "erp" / "master_data" / "stores.csv").write_text(MASTER, encoding="utf-8")
    write_report(root, relative, sheets)
    return root


def store_column(**overrides):
    column = {
        "metric": "net_sales",
        "unit": "money",
        "currency_from": "Ccy",
        "vat_included": False,
        "gross_or_net": "net",
        "definition_quote": "Net sales = gross sales",
    }
    column.update(overrides)
    return column


def visitors_column(**overrides):
    column = {"metric": "traffic", "unit": "visitors", "definition_quote": "Traffic from door"}
    column.update(overrides)
    return column


FLASH_PATTERN = "retail_controlling/weekly_flash/*.xlsx"


def write_config(tmp_path, *, columns, summary=None, pattern=FLASH_PATTERN, status="approved"):
    config = tmp_path / "config"
    config.mkdir(exist_ok=True)
    raw = {
        "version": 1,
        "mapping_version": "test-v1",
        "status": status,
        "report_pattern": pattern,
        "calendar": "retail_fiscal_week",
        "sheets": {
            "Stores": {
                "header_row": 2,
                "first_row": 3,
                "key_column": "Store",
                "scope": "store",
                "lfl_column": "LFL",
                "columns": columns,
            }
        },
    }
    if summary is not None:
        raw["sheets"]["Summary"] = {
            "header_row": 4,
            "first_row": 5,
            "key_column": "Region",
            "scope": "region",
            "columns": summary,
        }
    (config / "report_metrics.yaml").write_text(yaml.safe_dump(raw, allow_unicode=True))
    return config


def test_registry_discovers_ing_metrics():
    assert MANIFEST in Registry.discover().manifests()


def test_the_real_flash_publishes_every_mapped_metric(tmp_path):
    request = spec(all_stores())
    result, facts = facts_of(tmp_path, request)
    assert result.status == "succeeded"
    assert result.counts["documents"] == 1
    assert result.counts["facts"] == len(facts)
    keys = by_key(facts)
    milan = keys[("MIL01", "net_sales")]
    assert milan.value == 589265.0
    assert milan.unit == "money"
    assert milan.currency == "EUR"
    assert milan.vat_included is False
    assert milan.gross_or_net == "net"
    assert milan.locator == "Stores!F3"
    assert milan.mapping_version == "weekly-flash-v1"
    assert milan.calendar == "retail_fiscal_week"
    assert milan.scope == {
        "level": "store",
        "store_id": "MIL01",
        "region": "EMEA",
        "country_code": "IT",
    }
    assert milan.document_id == "doc-flash-w42-2026"
    assert milan.period_start.isoformat() == "2026-10-11T22:00:00+00:00"
    assert milan.period_end.isoformat() == "2026-10-18T21:59:59+00:00"
    assert milan.available_at.isoformat() == "2026-10-19T22:00:00+00:00"  # issue date, Milan
    assert milan.comparability_status == "comparable"
    assert milan.definition.startswith("Net sales = gross sales - discounts - returns")
    assert milan.fx_basis is None  # local currency: no FX basis to state
    assert keys[("MIL01", "traffic")].value == 3957.0
    assert keys[("MIL01", "traffic")].unit == "visitors"
    assert milan.source_refs[0].locator == milan.locator
    assert (
        milan.source_refs[0].uri
        == "dataset://reports/retail_controlling/weekly_flash/Flash_W42_2026.xlsx"
    )
    assert milan.input_refs[0].kind == "SourceDocument"


def test_the_local_column_keeps_the_rows_own_currency(tmp_path):
    request = spec(all_stores())
    _, facts = facts_of(tmp_path, request)
    keys = by_key(facts)
    geneva = keys[("GVA01", "net_sales")]
    assert (geneva.value, geneva.currency) == (169165.0, "CHF")
    reported_eur = keys[("GVA01", "net_sales_eur")]
    assert (reported_eur.value, reported_eur.currency) == (183416.4588528678, "EUR")
    assert reported_eur.fx_basis.startswith("Reported EUR at average October 2026 rates")


def test_the_lfl_column_is_a_percent_and_only_for_lfl_stores(tmp_path):
    request = spec(all_stores())
    _, facts = facts_of(tmp_path, request)
    keys = by_key(facts)
    milan = keys[("MIL01", "net_sales_growth_lfl")]
    assert milan.unit == "percent"
    assert milan.value == pytest.approx(17.47707336523125)
    assert milan.definition.startswith("LFL = stores trading at least 5 days")
    # A fall is a fact too: Roma is -40%, not a dropped negative count.
    assert keys[("ROM01", "net_sales_growth_lfl")].value < 0
    # LFL = N rows store no growth at all: no fact, and certainly not zero.
    assert ("COU01", "net_sales_growth_lfl") not in keys
    assert all(fact.value != 0 or fact.metric != "net_sales_growth_lfl" for fact in facts)


def test_region_totals_reconcile_with_the_store_rows(tmp_path):
    request = spec(all_stores())
    _, facts = facts_of(tmp_path, request)
    keys = by_key(facts)
    for region in ("EMEA", "AMERICAS", "APAC"):
        stores = sum(
            fact.value
            for fact in facts
            if fact.metric == "net_sales_eur"
            and fact.scope.get("level") == "store"
            and fact.scope.get("region") == region
        )
        assert keys[(region, "net_sales_eur")].value == pytest.approx(stores, abs=0.01)


def test_a_region_total_is_not_published_for_a_narrow_run(tmp_path):
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request)
    keys = by_key(facts)
    assert set(keys) == {
        ("MIL01", "net_sales"),
        ("MIL01", "net_sales_eur"),
        ("MIL01", "net_sales_growth_lfl"),
        ("MIL01", "traffic"),
    }
    assert result.counts["rows_outside_run"] == 41  # 38 other stores + the three regions
    assert any("EMEA is outside this run's stores" in note for note in result.abstentions)


def test_an_unknown_region_is_an_abstention_not_a_fact(tmp_path):
    request = spec(all_stores())
    result, facts = facts_of(tmp_path, request)
    assert all(fact.scope.get("region") != "TOTAL RETAIL" for fact in facts)
    assert any("TOTAL RETAIL" in note for note in result.abstentions)


def test_zero_is_a_real_number_and_a_missing_cell_is_not(tmp_path):
    request = spec(all_stores())
    _, facts = facts_of(tmp_path, request)
    keys = by_key(facts)
    courmayeur = keys[("COU01", "net_sales")]
    assert courmayeur.value == 0.0
    assert courmayeur.comparability_status == "comparable"


def test_a_text_cell_is_not_a_zero(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": COVER,
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "EUR", "n.a.", 100, 10, "Y", 0.1]
            ),
        },
    )
    config = write_config(
        tmp_path,
        columns={
            "Net Sales LC": store_column(),
            "Net Sales EUR": store_column(
                metric="net_sales_eur",
                currency="EUR",
                currency_from=None,
                fx_basis_quote="Reported EUR at average October",
            ),
        },
    )
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    # The textual cell is dropped; the number next to it is still a fact, not 0.
    assert [(fact.metric, fact.value) for fact in facts] == [("net_sales_eur", 100.0)]
    assert result.counts["cells_not_numeric"] == 1
    assert any("'n.a.' is not a number" in note for note in result.abstentions)


def test_a_textual_number_is_not_trusted(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": COVER,
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "EUR", "589265", 1, 1, "Y", 0.1]
            ),
        },
    )
    config = write_config(tmp_path, columns={"Net Sales LC": store_column()})
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    assert facts == []
    assert result.counts["cells_not_numeric"] == 1


def test_an_unparseable_currency_drops_the_fact_instead_of_defaulting(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": COVER,
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "euro", 10, 10, 1, "Y", 0.1]
            ),
        },
    )
    config = write_config(tmp_path, columns={"Net Sales LC": store_column()})
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    assert facts == []
    assert result.counts["cells_without_currency"] == 1


def test_a_definition_missing_from_the_report_blocks_comparability(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": [line for line in COVER if not line or "Net sales =" not in line[0]],
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "EUR", 10, 10, 1, "Y", 0.1]
            ),
        },
    )
    config = write_config(tmp_path, columns={"Net Sales LC": store_column()})
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    fact = facts[0]
    assert fact.value == 10.0
    assert fact.definition is None
    assert fact.comparability_status == "not_comparable"
    assert result.counts["definitions_missing"] == 1


def test_a_missing_fx_basis_blocks_comparability_of_the_eur_column(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": [line for line in COVER if not line or "Reported EUR" not in line[0]],
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "EUR", 10, 10, 1, "Y", 0.1]
            ),
        },
    )
    config = write_config(
        tmp_path,
        columns={
            "Net Sales EUR": store_column(
                metric="net_sales_eur",
                currency="EUR",
                currency_from=None,
                fx_basis_quote="Reported EUR at average October",
            )
        },
    )
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    assert facts[0].comparability_status == "not_comparable"
    assert facts[0].fx_basis is None
    assert result.counts["fx_basis_missing"] == 1


def test_a_non_lfl_store_is_not_comparable_even_with_a_number(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": COVER,
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "EUR", 10, 10, 1, "N", 0.1]
            ),
        },
    )
    config = write_config(
        tmp_path,
        columns={
            "\u0394% LFL cFX": {
                "metric": "net_sales_growth_lfl",
                "unit": "percent",
                "scale": 100,
                "requires_lfl": True,
                "definition_quote": "LFL = stores trading at least 5 days",
            }
        },
    )
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    assert facts[0].value == pytest.approx(10.0)
    assert facts[0].comparability_status == "not_comparable"
    assert result.counts["lfl_not_comparable"] == 1


def test_scale_declares_eur_thousands(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": COVER,
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "EUR", 10, 10, 1, "Y", 0.1]
            ),
            "Summary": summary_sheet(["EMEA", 4520.929947622046]),
        },
    )
    config = write_config(
        tmp_path,
        columns={"Net Sales LC": store_column()},
        summary={
            "Net Sales Wk (EUR k)": store_column(
                metric="net_sales_eur",
                currency="EUR",
                currency_from=None,
                scale=1000,
            )
        },
    )
    request = spec(all_stores())
    _, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    region = next(fact for fact in facts if fact.scope["level"] == "region")
    assert region.value == pytest.approx(4520929.947622046)
    assert region.scope == {"level": "region", "region": "EMEA"}


def test_a_column_missing_from_the_sheet_keeps_the_other_facts(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": COVER,
            "Stores": [
                ["Store detail"],
                ["Store", "Ccy", "Net Sales LC"],
                ["Milano Montenapoleone", "EUR", 10],
            ],
        },
    )
    config = write_config(
        tmp_path,
        columns={"Net Sales LC": store_column(), "Traffic": visitors_column()},
    )
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    assert [fact.metric for fact in facts] == ["net_sales"]
    assert any("column 'Traffic' is missing" in note for note in result.abstentions)


def test_a_negative_physical_count_is_dropped(tmp_path):
    sources = write_sources(
        tmp_path,
        sheets={
            "Cover": COVER,
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "EUR", 10, 10, -3, "Y", -0.1]
            ),
        },
    )
    config = write_config(
        tmp_path,
        columns={
            "Net Sales LC": store_column(),
            "Traffic": visitors_column(),
        },
    )
    request = spec(("MIL01",))
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    assert [(fact.metric, fact.value) for fact in facts] == [("net_sales", 10.0)]
    assert result.counts["cells_negative"] == 1


def test_an_unmapped_document_is_abstained(tmp_path):
    sources = write_sources(
        tmp_path,
        relative="other/Board_pack.xlsx",
        sheets={
            "Cover": COVER,
            "Stores": stores_sheet(
                ["1", "Milano Montenapoleone", "EMEA", "Italy", "EUR", 10, 10, 1, "Y", 0.1]
            ),
        },
    )
    config = write_config(tmp_path, columns={"Net Sales LC": store_column()})
    request = spec(("MIL01",), report="other/Board_pack.xlsx")
    result, facts = facts_of(tmp_path, request, sources_dir=sources, config_dir=config)
    assert facts == []
    assert result.counts["documents_unmapped"] == 1


def test_fact_ids_are_stable_for_the_same_content(tmp_path):
    request = spec(all_stores())
    _, facts = facts_of(tmp_path, request)
    (tmp_path / "again").mkdir()
    again = facts_of(tmp_path / "again", request.model_copy(update={"run_id": "metrics-again"}))[1]
    assert sorted(fact.id for fact in facts) == sorted(fact.id for fact in again)
    assert len({fact.id for fact in facts}) == len(facts)


def test_an_unapproved_mapping_fails_the_run(tmp_path):
    config = write_config(tmp_path, columns={"Net Sales LC": store_column()}, status="draft")
    request = spec(("MIL01",))
    _, statuses = execute(tmp_path, request, config_dir=config)
    result = statuses[MANIFEST.id].result
    assert statuses[MANIFEST.id].status == "failed"
    assert "not approved" in result.errors[0]


@pytest.mark.parametrize(
    "column, message",
    [
        ({"unit": "money", "currency": None, "currency_from": None}, "exactly one of currency"),
        ({"unit": "visitors", "currency": "EUR"}, "only money columns carry a currency"),
        ({"scale": 0}, "scale must be a positive"),
        ({"vat_included": "no"}, "vat_included must be true or false"),
        ({"gross_or_net": "turnover"}, "gross_or_net must be"),
        ({"extra": 1}, "unknown key"),
        ({"definition_quote": ""}, "definition_quote must be"),
    ],
)
def test_a_bad_column_is_refused(tmp_path, column, message):
    path = write_config(tmp_path, columns={"Net Sales LC": store_column(**column)})
    with pytest.raises(metrics.MappingError, match=message):
        metrics.load_mapping(path / "report_metrics.yaml")


def test_two_columns_cannot_share_one_metric(tmp_path):
    path = write_config(
        tmp_path,
        columns={"Net Sales LC": store_column(), "Net Sales EUR": store_column()},
    )
    with pytest.raises(metrics.MappingError, match="same metric"):
        metrics.load_mapping(path / "report_metrics.yaml")


def test_the_mapping_cannot_escape_the_config_dir(tmp_path):
    with pytest.raises(metrics.MappingError, match="under config"):
        metrics.mapping_path(tmp_path, "../report_metrics.yaml")
