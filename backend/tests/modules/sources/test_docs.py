"""src.docs: the real W42 flash, and synthetic workbooks for the edges."""

import re
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from jsonschema.exceptions import ValidationError
from openpyxl import Workbook

from tesseract.core.contracts import ModuleSnapshot, RunSpec, Segment, SourceDocument
from tesseract.core.db import connect
from tesseract.core.orchestrator import run
from tesseract.core.registry import RegisteredModule, Registry
from tesseract.core.run_repo import repo_session
from tesseract.modules.sources import docs
from tesseract.modules.sources.docs import MANIFEST, parse_report, report_path

ROOT = Path(__file__).resolve().parents[4]
SOURCES = ROOT / "generated" / "sources"
FLASHES = "retail_controlling/weekly_flash"
AS_OF = datetime(2026, 10, 22, 20, 0, tzinfo=UTC)
LOCATOR = re.compile(r"^[^!]+(![A-Z]+[0-9]+(?::[A-Z]+[0-9]+)?)?$")


def flash(name):
    return f"{FLASHES}/{name}"


def spec(as_of=AS_OF, params=None, run_id="docs-run"):
    return RunSpec(
        run_id=run_id,
        scenario_id="historical",
        as_of=as_of,
        stores=("MIL01",),
        source_manifest_hash="sources",
        config_hash="config",
        code_version="test",
        mode="live",
        modules=(ModuleSnapshot(manifest=MANIFEST, params=params or {}),),
        prompt_versions={},
        model_versions={},
        authorized_fixture_refs=(),
    )


def execute(tmp_path, request, sources_dir=SOURCES):
    engine = connect(f"sqlite:///{tmp_path / 'docs.db'}")
    module = RegisteredModule(MANIFEST, lambda ctx: docs.run(ctx, sources_dir=sources_dir))
    registry = Registry((module,))
    statuses = run(request, engine=engine, registry=registry, clients={})
    return engine, statuses[0]


def parsed(relative=docs.DEFAULT_REPORT, sources_dir=SOURCES):
    return parse_report(report_path(sources_dir, relative), relative=relative)


def workbook(tmp_path, sheets, name="synthetic.xlsx"):
    """``sheets`` maps a title to rows of cell values; rows are written left to right."""
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row_index, values in enumerate(rows, start=1):
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row_index, column=column, value=value)
    path = tmp_path / "reports" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path.parent.parent


def test_registry_discovers_src_docs():
    assert MANIFEST in Registry.discover().manifests()


def test_cover_facts_come_from_the_workbook():
    report = parsed()
    assert report.issue_date.isoformat() == "2026-10-20"
    assert (report.period_start.isoformat(), report.period_end.isoformat()) == (
        "2026-10-12",
        "2026-10-18",
    )
    assert report.fiscal_week == 42
    assert report.version == "1.0"
    assert report.sheets == ("Cover", "Summary", "Stores", "Country", "Daily", "Top styles", "FX")
    assert "issued 20/10/2026" in report.issued_line


def test_document_is_available_only_from_the_issue_date():
    document, _ = docs.build_document(parsed())
    # The flash states a date, not a time: the day starts in Milan (00:00 CEST = 22:00 UTC before).
    assert document.available_at == datetime(2026, 10, 20, 0, 0, tzinfo=docs.REPORT_TIMEZONE)
    assert document.available_at == datetime(2026, 10, 19, 22, 0, tzinfo=UTC)
    assert document.availability_basis == "report_cover_issue_date_not_before_period_end"
    assert document.event_start == datetime(2026, 10, 12, 0, 0, tzinfo=docs.REPORT_TIMEZONE)
    assert document.event_start == datetime(2026, 10, 11, 22, 0, tzinfo=UTC)
    assert document.event_end == datetime(2026, 10, 18, 23, 59, 59, tzinfo=docs.REPORT_TIMEZONE)


def test_segments_are_the_rows_with_cell_locators():
    report = parsed()
    document, segments = docs.build_document(report)
    assert document.source_type == "report"
    assert document.connector == "src.docs"
    assert document.origin == "retail_controlling"
    assert document.language == "en"
    assert document.business_date.isoformat() == "2026-10-18"
    assert document.meta["issue_date"] == "2026-10-20"
    assert document.meta["text"] == docs.TEXT_BASIS
    assert len(segments) == len(report.rows)
    assert document.text_version == docs.TEXT_VERSION
    for index, segment in enumerate(segments):
        segment.validate_against(document)
        assert segment.idx == index
        assert segment.id == f"{document.id}-seg{index}"
        assert LOCATOR.match(segment.locator or "")
        assert segment.locator.split("!")[0] in document.meta["sheets"]
        assert segment.source_refs[0].locator == segment.locator
    # Document text is exactly the rows joined by one newline, no hidden separator.
    assert document.text == "\n".join(segment.text for segment in segments)
    assert document.text.count("\n") + 1 == len(segments)  # one row is one line
    assert segments[-1].char_end == len(document.text)


def test_the_slice_sheets_are_readable_with_provenance():
    document, segments = docs.build_document(parsed())
    by_locator = {segment.locator: segment for segment in segments}
    definitions = by_locator["Cover!B9"]
    assert definitions.text.startswith("Net sales = gross sales - discounts - returns")
    assert "'" not in definitions.locator  # locators address global cells, not local ranges
    milan = by_locator["Stores!A3:U3"]
    assert milan.text.startswith("1001 | Milano Montenapoleone | EMEA | Italy | EUR | 589265")
    assert "3957" in milan.text  # traffic, the number behind the S3 storyline
    cardigan = next(s for s in segments if s.text.startswith("14 | VLOGO CARDIGAN | Knitwear"))
    assert cardigan.locator == "Top styles!A16:F16"
    assert "295448.2923937891" in cardigan.text
    # The quote a later module will cite is an exact slice of the canonical text.
    assert document.text[cardigan.char_start : cardigan.char_end] == cardigan.text


def test_run_writes_the_report_at_as_of(tmp_path):
    request = spec()
    engine, status = execute(tmp_path, request)
    assert status.result.status == "succeeded"
    assert status.result.counts["documents"] == 1
    assert status.result.counts["sheets"] == 7
    with repo_session(engine, request) as repo:
        documents = repo.list(SourceDocument)
        segments = repo.list(Segment)
        assert [document.id for document in documents] == ["doc-flash-w42-2026"]
        assert len(segments) == status.result.counts["segments"] == len(parsed().rows)
        assert {segment.document_id for segment in segments} == {documents[0].id}


def test_report_is_not_read_before_its_issue_date(tmp_path):
    # 19/10 23:59 in Milan: the flash issued on 20/10 does not exist yet.
    as_of = datetime(2026, 10, 19, 21, 59, tzinfo=UTC)
    request = spec(as_of=as_of)
    engine, status = execute(tmp_path, request)
    assert status.result.status == "succeeded"
    assert status.result.counts == {"documents": 0, "segments": 0, "sheets": 0, "not_available": 1}
    assert status.result.abstentions == (
        f"dataset://reports/{docs.DEFAULT_REPORT}: not issued at as_of",
    )
    with repo_session(engine, request) as repo:
        assert repo.list(SourceDocument) == []


def test_a_correction_is_a_separate_document():
    first = parsed(flash("Flash_W40_2026_v1.xlsx"))
    second = parsed(flash("Flash_W40_2026_v2_corrected.xlsx"))
    before, _ = docs.build_document(first)
    after, _ = docs.build_document(second)
    assert (before.id, after.id) == ("doc-flash-w40-2026-v1", "doc-flash-w40-2026-v2-corrected")
    assert before.content_hash != after.content_hash
    assert (before.meta["version"], after.meta["version"]) == ("1.0", "2.0")
    assert before.event_end == after.event_end  # same week, corrected figures
    assert before.text != after.text


def test_an_undated_workbook_is_skipped_not_guessed(tmp_path):
    sources = workbook(tmp_path, {"Sheet1": [["Metric", "Value"], ["Sales", 10]]})
    request = spec(params={"report": "synthetic.xlsx"})
    engine, status = execute(tmp_path, request, sources_dir=sources)
    assert status.result.status == "succeeded"
    assert status.result.counts == {
        "documents": 0,
        "segments": 0,
        "sheets": 0,
        "reports_skipped": 1,
    }
    assert "no 'issued <date>' line" in status.result.abstentions[0]
    with repo_session(engine, request) as repo:
        assert repo.list(SourceDocument) == []


def test_row_text_keeps_columns_and_cell_types(tmp_path):
    sources = workbook(
        tmp_path,
        {
            "Cover": [
                ["WEEKLY SALES FLASH - WEEK 7 / 2026"],
                ["Fiscal week 7: Mon 09/02/2026 - Sun 15/02/2026"],
                ["Version 1.0 - issued 16/02/2026"],
            ],
            "Data": [
                ["Plant", "Store", "Units", "Note"],
                [1001, "Milano", 0, "n.a."],
                [1002, None, 1.5, True],
                [None, None, None, None],
            ],
            "Empty": [],
        },
    )
    document, segments = docs.build_document(parsed("synthetic.xlsx", sources_dir=sources))
    by_locator = {segment.locator: segment.text for segment in segments}
    assert by_locator["Data!A2:D2"] == "1001 | Milano | 0 | n.a."
    assert by_locator["Data!A3:D3"] == "1002 |  | 1.5 | TRUE"  # empty cell keeps its column
    assert "Data!A4:D4" not in by_locator  # an empty row is not a segment
    assert by_locator["Empty"] == "[Sheet: Empty]"  # an empty sheet still says it was there
    assert document.meta["version"] == "1.0"


def test_a_path_outside_the_reports_folder_is_an_error(tmp_path):
    reports = ("../store_emails/MIL01/20261009_1035_MIL01.eml", "/etc/hosts", "nope.xlsx")
    for index, relative in enumerate(reports):
        request = spec(params={"report": relative}, run_id=f"docs-run-{index}")
        _, status = execute(tmp_path, request)
        assert status.result.status == "degraded"
        assert status.result.errors
        assert status.result.counts["documents"] == 0


def test_a_csv_report_is_not_read_here(tmp_path):
    request = spec(params={"report": f"{FLASHES}/export_classifica_negozi_sett42.csv"})
    _, status = execute(tmp_path, request)
    assert status.result.status == "degraded"
    assert "only .xlsx reports are read here" in status.result.errors[0]


def test_params_must_name_a_report_as_text():
    registry = Registry((RegisteredModule(MANIFEST, docs.run),))
    with pytest.raises(ValidationError):
        registry.resolve(ModuleSnapshot(manifest=MANIFEST, params={"report": 3}))
    snapshot = ModuleSnapshot(manifest=MANIFEST, params={"report": flash("Flash_W41_2026.xlsx")})
    assert registry.resolve(snapshot).manifest is MANIFEST


def test_selected_report_param_wins_over_the_default(tmp_path):
    request = spec(params={"report": flash("Flash_W41_2026.xlsx")})
    engine, status = execute(tmp_path, request)
    assert status.result.counts["documents"] == 1
    with repo_session(engine, request) as repo:
        assert [document.id for document in repo.list(SourceDocument)] == ["doc-flash-w41-2026"]


def test_a_single_cell_row_gets_a_single_cell_locator():
    document, segments = docs.build_document(parsed())
    spans = {segment.locator: segment for segment in segments}
    assert spans["Cover!B2"].text == "VALENTINO | Retail Controlling"
    assert document.text[spans["Cover!B2"].char_start : spans["Cover!B2"].char_end] == (
        "VALENTINO | Retail Controlling"
    )


def test_availability_is_not_before_the_period_end(tmp_path):
    # A report issued on the 10th for the week ending the 15th still covers days after its issue.
    sources = workbook(
        tmp_path,
        {
            "Cover": [
                ["Fiscal week 7: Mon 09/02/2026 - Sun 15/02/2026"],
                ["Version 1.0 - issued 10/02/2026"],
                ["Sales", 10],
            ]
        },
    )
    report = parsed("synthetic.xlsx", sources_dir=sources)
    assert report.issue_date.isoformat() == "2026-02-10"
    assert report.available_at == datetime(2026, 2, 15, 0, 0, tzinfo=docs.REPORT_TIMEZONE)


def test_a_run_on_the_issue_date_itself_reads_the_report(tmp_path):
    # available_at is inclusive: 20/10 00:00 in Milan is exactly when it starts to count.
    as_of = datetime(2026, 10, 19, 22, 0, tzinfo=UTC)
    _, status = execute(tmp_path, spec(as_of=as_of))
    assert status.result.counts["documents"] == 1


def test_cell_line_breaks_become_spaces(tmp_path):
    sources = workbook(
        tmp_path,
        {
            "Cover": [
                ["Fiscal week 7: Mon 09/02/2026 - Sun 15/02/2026"],
                ["Version 1.0 - issued 16/02/2026"],
            ],
            "Data": [["Region", "Net Sales\nWk (EUR k)"], ["EMEA", 4520.5]],
        },
    )
    document, segments = docs.build_document(parsed("synthetic.xlsx", sources_dir=sources))
    assert {segment.locator: segment.text for segment in segments}["Data!A1:B1"] == (
        "Region | Net Sales Wk (EUR k)"
    )
    assert document.text.count("\n") + 1 == len(segments)


def test_a_broken_workbook_degrades_the_module_not_the_run(tmp_path):
    # openpyxl raises KeyError for a zip without the manifest, not BadZipFile.
    broken = tmp_path / "reports" / "broken.xlsx"
    broken.parent.mkdir(parents=True)
    with ZipFile(broken, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    request = spec(params={"report": "broken.xlsx"})
    _, status = execute(tmp_path, request, sources_dir=tmp_path)
    assert status.result.status == "degraded"
    assert status.result.counts["documents"] == 0
    assert "unreadable workbook (KeyError" in status.result.errors[0]


def test_a_report_of_another_family_says_what_is_missing(tmp_path):
    # The merchandising sell-through states "cut-off 18/10": a period end, not an issue date.
    request = spec(params={"report": "merchandising/FW26_NewIn_SellThrough_wk42_final.xlsx"})
    _, status = execute(tmp_path, request)
    assert status.result.status == "succeeded"
    assert status.result.counts == {
        "documents": 0,
        "segments": 0,
        "sheets": 0,
        "reports_skipped": 1,
    }
    assert "no 'issued <date>' line on any sheet" in status.result.abstentions[0]


def test_the_caller_declares_the_report_language(tmp_path):
    request = spec(params={"language": "it"})
    engine, status = execute(tmp_path, request)
    assert status.result.counts["documents"] == 1
    with repo_session(engine, request) as repo:
        assert [document.language for document in repo.list(SourceDocument)] == ["it"]


def test_a_language_that_is_not_a_tag_is_rejected():
    registry = Registry((RegisteredModule(MANIFEST, docs.run),))
    with pytest.raises(ValidationError):
        registry.resolve(ModuleSnapshot(manifest=MANIFEST, params={"language": "english"}))
