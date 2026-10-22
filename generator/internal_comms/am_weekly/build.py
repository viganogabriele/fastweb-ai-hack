"""Wrap hand-written area-manager recaps (msgs_*.py) into .eml / .docx.
Run: .venv/bin/python generator/internal_comms/am_weekly/build.py

Message dict keys:
  key, folder, file (no ext), frm, to (list), cc (list|None), subj, sent 'YYYY-MM-DD HH:MM', tz, body
  optional: reply_to (key of parent; adds In-Reply-To and a quoted parent), quote (bool, default True for replies),
            html (str), docx (dict(name, title, paras[list[str]], table[list[list[str]]] | None)) -> attached,
            standalone_docx (same dict, written as a separate .docx file and no eml)
"""
import importlib
import io
import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
from common import OUT, make_eml, write_eml, check_texts  # noqa: E402
from docx import Document
from docx.shared import Pt

D = OUT / "area_manager_weekly"
MODULES = ["msgs_galli", "msgs_moreau", "msgs_wagner", "msgs_mcleod", "msgs_haddad", "msgs_mitchell",
           "msgs_suzuki", "msgs_yoon", "msgs_zhao", "msgs_lau", "msgs_lim"]


def docx_bytes(spec):
    doc = Document()
    doc.styles["Normal"].font.name = spec.get("font", "Calibri")
    doc.styles["Normal"].font.size = Pt(10)
    p = doc.add_paragraph()
    r = p.add_run(spec["title"])
    r.bold = True
    r.font.size = Pt(13)
    for para in spec.get("paras_before", []):
        doc.add_paragraph(para)
    if spec.get("table"):
        t = doc.add_table(rows=len(spec["table"]), cols=len(spec["table"][0]))
        t.style = "Table Grid"
        for i, row in enumerate(spec["table"]):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
    for para in spec.get("paras", []):
        doc.add_paragraph(para)
    if spec.get("author"):
        doc.core_properties.author = spec["author"]
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def quote(parent, msg_obj, lang):
    hdr = {
        "it": "-----Messaggio originale-----\nDa: {f}\nInviato: {d}\nA: {t}\nOggetto: {s}\n\n",
        "fr": "Le {d}, {f} a écrit :\n",
        "de": "Von: {f}\nGesendet: {d}\nAn: {t}\nBetreff: {s}\n\n",
        "ja": "{d} {f}:\n",
        "ko": "-----Original Message-----\nFrom: {f}\nSent: {d}\nTo: {t}\nSubject: {s}\n\n",
        "zh": "------------------ 原始邮件 ------------------\n发件人: {f}\n发送时间: {d}\n收件人: {t}\n主题: {s}\n\n",
    }.get(lang, "From: {f}\nSent: {d}\nTo: {t}\nSubject: {s}\n\n")
    h = hdr.format(f=parent["frm"], d=msg_obj["Date"], t=", ".join(parent["to"]), s=parent["subj"])
    body = "\n".join(("> " + ln) if lang in ("fr", "ja") else ln for ln in parent["body"].rstrip().splitlines())
    return "\n\n" + h + body + "\n"


def main():
    allmsgs, objs, texts = {}, {}, {}
    n = 0
    for modname in MODULES:
        try:
            mod = importlib.import_module(modname)
        except ModuleNotFoundError:
            print("missing", modname)
            continue
        for m in mod.MSGS:
            allmsgs[m["key"]] = m
            texts[m["key"]] = m["body"] + " ".join((m.get("docx") or m.get("standalone_docx") or {}).get("paras", []))
            folder = D / m["folder"]
            folder.mkdir(parents=True, exist_ok=True)
            if m.get("standalone_docx"):
                (folder / (m["file"] + ".docx")).write_bytes(docx_bytes(m["standalone_docx"]))
                n += 1
                continue
            body = m["body"]
            irt = None
            if m.get("reply_to"):
                parent = allmsgs[m["reply_to"]]
                irt = objs[m["reply_to"]]["Message-ID"]
                if m.get("quote", True):
                    body = body + quote(parent, objs[m["reply_to"]], getattr(mod, "LANG", "en"))
            att = []
            if m.get("docx"):
                att.append((m["docx"]["name"], docx_bytes(m["docx"]), "application",
                            "vnd.openxmlformats-officedocument.wordprocessingml.document"))
            e = make_eml(m["frm"], m["to"], m["subj"], body, sent=m["sent"], tz=m["tz"], cc=m.get("cc"),
                         html=m.get("html"), attachments=att, in_reply_to=irt, msgid_seed=m["key"])
            objs[m["key"]] = e
            write_eml(e, folder / (m["file"] + ".eml"))
            n += 1
    print("files:", n)
    check_texts(texts)


if __name__ == "__main__":
    main()
