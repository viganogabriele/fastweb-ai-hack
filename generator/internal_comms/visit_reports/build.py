"""Render store visit reports (content in content.py). Run: .venv/bin/python generator/internal_comms/visit_reports/build.py"""
import pathlib
import subprocess
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
from common import OUT, check_texts  # noqa: E402
import content as C  # noqa: E402

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

D = OUT / "store_visit_reports"
D.mkdir(parents=True, exist_ok=True)
TEXTS = {}

SECTION_NAMES = {
    "it": C.V3_SECTIONS,
    "fr": ["Extérieur / vitrines", "Entrée et accessibilité", "Propreté et rangement", "VM intérieur", "Disponibilité produit",
           "Équipe et tenue", "Accueil et vente", "Clienteling", "Réserve / stock", "Sécurité"],
    "de": ["Außen / Schaufenster", "Eingang & Zugang", "Ordnung & Sauberkeit", "VM innen", "Warenverfügbarkeit",
           "Team & Erscheinungsbild", "Begrüßung & Verkauf", "Clienteling", "Lager / Backoffice", "Sicherheit"],
    "en": ["Exterior / windows", "Entrance & access", "Housekeeping", "Interior VM", "Product availability",
           "Team & grooming", "Welcome & selling", "Clienteling", "Back office / stockroom", "Safety"],
}
LBL = {
    "it": dict(store="Negozio", date="Data visita", visitor="Visitatore", att="Presenti", kpi="KPI", sec="Area", w="Peso", s="Voto (1-5)",
               notes="Commenti", act="Azioni", who="Chi", due="Entro"),
    "fr": dict(store="Boutique", date="Date", visitor="Visiteur", att="Présents", kpi="KPI", sec="Rubrique", w="Poids", s="Note (1-5)",
               notes="Commentaires", act="Actions", who="Qui", due="Échéance"),
    "de": dict(store="Filiale", date="Datum", visitor="Besucher", att="Anwesend", kpi="KPI", sec="Bereich", w="Gewicht", s="Note (1-5)",
               notes="Kommentare", act="Maßnahmen", who="Wer", due="Bis"),
    "en": dict(store="Store", date="Date", visitor="Visitor", att="Attendees", kpi="KPI", sec="Area", w="Weight", s="Score (1-5)",
               notes="Comments", act="Actions", who="Owner", due="Due"),
}


def to_pdf(path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(path.parent), str(path)],
                   check=True, capture_output=True)
    path.unlink()


def kv_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        t.cell(i, 0).text = k
        t.cell(i, 1).text = v
        t.cell(i, 0).paragraphs[0].runs[0].bold = True
    return t


def render_v3(r):
    L = LBL[r["lang"]]
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9.5)
    h = doc.add_paragraph()
    run = h.add_run(r["title"])
    run.bold = True
    run.font.size = Pt(14)
    kv_table(doc, [(L["store"], r["store"]), (L["date"], r["date"]), (L["visitor"], r["visitor"]), (L["att"], r["attendees"]), (L["kpi"], r["kpi"])])
    doc.add_paragraph()
    names = SECTION_NAMES[r["lang"]]
    t = doc.add_table(rows=len(names) + 2, cols=3)
    t.style = "Table Grid"
    for j, v in enumerate([L["sec"], L["w"], L["s"]]):
        t.cell(0, j).text = v
    tot_w = tot = 0
    for i, (n, w, s) in enumerate(zip(names, C.V3_WEIGHTS, r["scores"]), start=1):
        t.cell(i, 0).text = n
        t.cell(i, 1).text = f"{w}%"
        t.cell(i, 2).text = "n/a" if s is None else str(s)
        if s is not None:
            tot_w += w
            tot += w * s / 5
    t.cell(len(names) + 1, 0).text = r["total_label"]
    t.cell(len(names) + 1, 2).text = f"{round(100 * tot / tot_w)}%"
    doc.add_paragraph()
    doc.add_paragraph(L["notes"]).runs[0].bold = True
    for n in r["notes"]:
        doc.add_paragraph(n, style="List Bullet")
    doc.add_paragraph(L["act"]).runs[0].bold = True
    a = doc.add_table(rows=len(r["actions"]) + 1, cols=3)
    a.style = "Table Grid"
    for j, v in enumerate([L["act"], L["who"], L["due"]]):
        a.cell(0, j).text = v
    for i, row in enumerate(r["actions"], start=1):
        for j, v in enumerate(row):
            a.cell(i, j).text = v
    doc.sections[0].footer.paragraphs[0].text = "Store Visit Checklist v3 – Retail Excellence – internal"
    p = D / r["file"]
    doc.save(p)
    if r.get("to_pdf"):
        to_pdf(p)
    TEXTS[r["file"]] = " ".join(r["notes"]) + r["kpi"] + " ".join(x[0] for x in r["actions"])


def render_vm(r):
    ss = getSampleStyleSheet()
    st = [Paragraph("VM COMPLIANCE CHECK – STORE VISIT", ParagraphStyle("t", parent=ss["Title"], fontSize=15, alignment=0)),
          Paragraph(f"{r['store']} &nbsp;|&nbsp; {r['date']} &nbsp;|&nbsp; {r['visitor']}", ss["Normal"]), Spacer(1, 8)]
    data = [["Item", "Status", "Photo ref", "Notes"]] + [[a, b, c, Paragraph(d, ParagraphStyle("c", fontSize=7.5, leading=9))] for a, b, c, d in r["rows"]]
    t = Table(data, colWidths=[4.2 * cm, 1.4 * cm, 1.8 * cm, 9.6 * cm])
    style = [("FONT", (0, 0), (-1, -1), "Helvetica", 8), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8), ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
             ("VALIGN", (0, 0), (-1, -1), "TOP")]
    for i, row in enumerate(r["rows"], start=1):
        if row[1] == "KO":
            style.append(("TEXTCOLOR", (1, i), (1, i), colors.red))
    t.setStyle(TableStyle(style))
    st += [t, Spacer(1, 10), Paragraph("<b>Comment</b>", ss["Normal"]), Paragraph(r["comment"], ss["Normal"]), Spacer(1, 10),
           Paragraph("[photos W1-01 ... FR-02 uploaded to the VM SharePoint folder 'MIL01/2026-10']", ParagraphStyle("g", parent=ss["Normal"], textColor=colors.grey, fontSize=7))]
    SimpleDocTemplate(str(D / r["file"]), pagesize=A4, title="VM check MIL01").build(st)
    TEXTS[r["file"]] = " ".join(x[3] for x in r["rows"]) + r["comment"]


def render_jp(r):
    doc = Document()
    doc.styles["Normal"].font.name = "Noto Sans CJK JP"
    doc.styles["Normal"].font.size = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("店舗巡回レポート")
    run.bold = True
    run.font.size = Pt(15)
    kv_table(doc, r["head"])
    doc.add_paragraph("評価：◎ 大変良い ／ ○ 良い ／ △ 要改善 ／ × 不可")
    t = doc.add_table(rows=len(r["checks"]) + 1, cols=3)
    t.style = "Table Grid"
    for j, v in enumerate(["項目", "評価", "コメント"]):
        t.cell(0, j).text = v
    for i, row in enumerate(r["checks"], start=1):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    doc.add_paragraph("【所感】").runs[0].bold = True
    for line in r["comment"].split("\n"):
        doc.add_paragraph(line)
    doc.add_paragraph("【アクション】").runs[0].bold = True
    for a, w, d in r["actions"]:
        doc.add_paragraph(f"・{a}（担当：{w}／期限：{d}）")
    pth = D / r["file"]
    doc.save(pth)
    to_pdf(pth)
    TEXTS[r["file"]] = r["comment"] + "".join(c[2] for c in r["checks"])


def render_kr(r):
    doc = Document()
    doc.styles["Normal"].font.name = "Noto Sans CJK KR"
    doc.styles["Normal"].font.size = Pt(10)
    p = doc.add_paragraph()
    run = p.add_run("매장 방문 보고서")
    run.bold = True
    run.font.size = Pt(15)
    kv_table(doc, r["head"])
    doc.add_paragraph()
    t = doc.add_table(rows=len(r["score_rows"]) + 2, cols=3)
    t.style = "Table Grid"
    for j, v in enumerate(["항목", "점수", "만점"]):
        t.cell(0, j).text = v
    tot = 0
    for i, (a, s, m) in enumerate(r["score_rows"], start=1):
        t.cell(i, 0).text, t.cell(i, 1).text, t.cell(i, 2).text = a, str(s), str(m)
        tot += s
    n = len(r["score_rows"]) + 1
    t.cell(n, 0).text, t.cell(n, 1).text, t.cell(n, 2).text = "합계", str(tot), "100"
    doc.add_paragraph("■ 주요 내용").runs[0].bold = True
    for b in r["body"]:
        doc.add_paragraph(b)
    doc.add_paragraph("■ 후속 조치").runs[0].bold = True
    for a, w, d in r["actions"]:
        doc.add_paragraph(f"- {a} / 담당: {w} / 기한: {d}")
    doc.save(D / r["file"])
    TEXTS[r["file"]] = " ".join(r["body"])


def render_cn(r):
    doc = Document()
    doc.styles["Normal"].font.name = "Noto Sans CJK SC"
    doc.styles["Normal"].font.size = Pt(10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("门店巡店报告")
    run.bold = True
    run.font.size = Pt(15)
    kv_table(doc, r["head"])
    doc.add_paragraph()
    t = doc.add_table(rows=len(r["items"]) + 2, cols=3)
    t.style = "Table Grid"
    for j, v in enumerate(["检查项目", "得分（10分制）", "备注"]):
        t.cell(0, j).text = v
    tot = 0
    for i, (a, s, n) in enumerate(r["items"], start=1):
        t.cell(i, 0).text, t.cell(i, 1).text, t.cell(i, 2).text = a, str(s), n
        tot += s
    k = len(r["items"]) + 1
    t.cell(k, 0).text, t.cell(k, 1).text = "平均", f"{tot / len(r['items']):.1f}"
    doc.add_paragraph("巡店总结：").runs[0].bold = True
    doc.add_paragraph(r["summary"])
    doc.add_paragraph("跟进事项：").runs[0].bold = True
    for i, (a, w, d) in enumerate(r["actions"], start=1):
        doc.add_paragraph(f"{i}. {a}　负责人：{w}　时间：{d}")
    doc.save(D / r["file"])
    TEXTS[r["file"]] = r["summary"]


def render_simple(r, title):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(14)
    kv_table(doc, r["head"])
    doc.add_paragraph()
    kv_table(doc, r["grid"])
    doc.add_paragraph()
    for x in r["text"]:
        doc.add_paragraph(x)
    doc.add_paragraph("Follow-ups").runs[0].bold = True
    for a, w, d in r["actions"]:
        doc.add_paragraph(f"{a} – {w} – {d}", style="List Bullet")
    pth = D / r["file"]
    doc.save(pth)
    if r.get("to_pdf"):
        to_pdf(pth)
    TEXTS[r["file"]] = " ".join(r["text"])


def render_hk(r):
    ss = getSampleStyleSheet()
    st = []
    for txt, k in r["lines"]:
        if k == "t":
            st.append(Paragraph(txt, ss["Title"]))
        elif k == "h":
            st.append(Paragraph(txt, ss["Heading3"]))
        elif k == "b":
            st.append(Spacer(1, 8))
            st.append(Paragraph(f"<b>{txt}</b>", ss["Normal"]))
        else:
            st.append(Paragraph(txt, ss["Normal"]))
    SimpleDocTemplate(str(D / r["file"]), pagesize=A4, title="HKG01 visit").build(st)
    TEXTS[r["file"]] = " ".join(t for t, _ in r["lines"])


if __name__ == "__main__":
    for r in C.REPORTS:
        render_v3(r)
    render_vm(C.VM_MIL01)
    render_jp(C.JP_TYO01)
    render_kr(C.KR_SEL01)
    render_cn(C.CN_SHA01)
    render_simple(C.US_NYC02, "STORE WALK – AMERICAS RETAIL")
    render_simple(C.AU_SYD01, "Store visit notes")
    render_hk(C.HK_HKG01)
    check_texts(TEXTS)
    for p in sorted(D.iterdir()):
        print(p.name, p.stat().st_size)
