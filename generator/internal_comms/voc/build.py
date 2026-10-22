"""Build voice-of-customer files. Free text is hand-written in *_rows.py / ocr_cards.py; code only lays out and
injects documented data-quality noise. Run: .venv/bin/python generator/internal_comms/voc/build.py"""
import csv
import datetime as dt
import pathlib
import random
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
from common import OUT, check_texts  # noqa: E402
import clienteling_rows as CR  # noqa: E402
import feedback_rows as FR  # noqa: E402
import ocr_cards as OC  # noqa: E402
import openpyxl
from openpyxl.styles import Font, PatternFill

D = OUT / "voice_of_customer"
D.mkdir(parents=True, exist_ok=True)
rnd = random.Random(26)

PHONE = {"MIL": "+39 3** *** **{}", "ROM": "+39 3** *** **{}", "VEN": "+39 3** *** **{}", "Mil": "+39 3** *** **{}", "MI-": "+39 3** *** **{}",
         "100": "+39 3** *** **{}", "Rin": "+39 3** *** **{}", "Rom": "+39 3** *** **{}", "PAR": "+33 6 ** ** ** {}", "Gal": "+33 6 ** ** ** {}",
         "GVA": "+41 79 *** ** {}", "CAN": "+33 6 ** ** ** {}", "LON": "+44 7*** ***{}", "Har": "+44 7*** ***{}", "AMS": "+31 6 **** **{}",
         "CPH": "+45 ** ** ** {}", "MUN": "+49 1** *****{}", "ZRH": "+41 79 *** ** {}", "BER": "+49 1** *****{}", "MAD": "+34 6** *** *{}",
         "NYC": "+1 (***) ***-**{}", "TOR": "+1 (***) ***-**{}", "BHS": "+1 (***) ***-**{}", "TYO": "+81 90-****-**{}", "銀座": "+81 90-****-**{}",
         "OSA": "+81 80-****-**{}", "SEL": "+82 10-****-**{}", "청담": "+82 10-****-**{}", "현대": "+82 10-****-**{}", "BUS": "+82 10-****-**{}",
         "SHA": "+86 138 **** **{}", "BJS": "+86 139 **** **{}", "CDU": "+86 136 **** **{}", "HKG": "+852 **** **{}", "TPE": "+886 9** *** *{}",
         "SIN": "+65 9*** **{}", "SYD": "+61 4** *** *{}"}

# documented overrides (see answer key)
WRONG_STORE = {("2026-10-15", "Priya Raman"): "LON01"}                 # LON02 SA logged under LON01
WRONG_OUTCOME = {("2026-10-18", "박민수", "CL28913290"): "no answer"}    # note says client came, outcome says no answer


def build_clienteling():
    rows = sorted(CR.ROWS, key=lambda r: r[0])
    out = []
    n = 4810
    for r in rows:
        date, store, sa, cid, tier, ch, dirn, reason, prod, outcome, note = r
        store = WRONG_STORE.get((date, sa), store)
        outcome = WRONG_OUTCOME.get((date, sa, cid), outcome)
        n += rnd.choice([1, 1, 1, 2, 3])
        d = dt.date.fromisoformat(date)
        ds = d.isoformat()
        if rnd.random() < 0.06:
            ds = d.strftime("%d/%m/%Y")  # legacy tablet rows
        phone = PHONE.get(store[:3], "").format(f"{rnd.randint(10, 99)}") if cid else ""
        t = f"{rnd.randint(10, 20):02d}:{rnd.randint(0, 59):02d}"
        out.append([f"CLG-26-{n:05d}", store, sa, cid, tier, ch, ds, t, dirn, reason, prod, outcome, note, phone])
    # sync glitch: some rows exported twice
    for i in sorted(rnd.sample(range(len(out)), 5)):
        out.insert(i + 1, list(out[i]))
    hdr = ["log_id", "store", "sa_name", "client_id", "client_tier", "channel", "contact_date", "contact_time", "direction",
           "reason", "product_mentioned", "outcome", "notes", "phone_masked"]
    with open(D / "clienteling_log_export_2026-09-01_2026-10-21.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(hdr)
        w.writerows(out)
    return {f"cl{i}": r[10] for i, r in enumerate(CR.ROWS) if len(r[10].split()) > 8}


def build_feedback():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "responses"
    hdr = ["response_id", "submitted_at", "store", "device", "survey_version", "language", "nps_0_10", "q1_product_availability",
           "q2_staff", "q3_fitting_advice", "q4_wait_time", "product_asked_not_available", "comment"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D2B44")
    rows = sorted(FR.ROWS, key=lambda r: r[0])
    seq = 10233
    dup_idx = set(rnd.sample(range(len(rows)), 3))
    for i, r in enumerate(rows):
        ts, store, lang, nps, a, s, fi, w, unav, com = r
        seq += rnd.choice([1, 1, 2, 4])
        ver = "v2" if ts >= "2026-10-01" else "v1"
        dev = f"{store}-TAB{1 if rnd.random() < 0.7 else 2}"
        line = [f"R{seq}", ts if len(ts) > 10 else ts, store, dev, ver, lang, nps, a, s, fi, w, unav or None, com or None]
        ws.append(line)
        if i in dup_idx:  # double submit (button tapped twice)
            l2 = list(line)
            l2[0] = f"R{seq + 1}"
            seq += 1
            ws.append(l2)
    for col, wdt in zip("ABCDEFGHIJKLM", [10, 17, 8, 12, 8, 6, 8, 10, 8, 10, 8, 26, 60]):
        ws.column_dimensions[col].width = wdt
    rd = wb.create_sheet("README")
    for line in [
        "In-store feedback tablet export (Medallia-lite). Exported 22/10/2026 07:05 CET by CRM EMEA (G. Sala).",
        "nps_0_10: 'How likely are you to recommend us' 0-10.",
        "q1..q4: satisfaction 1-5 (5 = excellent). Empty = skipped.",
        "v1 -> v2 (01/10/2026): q4 changed from 'waiting time to be served' to 'waiting time (fitting room / cash desk / tax free)'.",
        "Store tablets in APAC were activated late (SEL01/TYO01 from 01/10, others partial). Not all stores have tablets.",
        "Known issue: double tap on 'Send' creates two responses a few seconds apart.",
    ]:
        rd.append([line])
    rd.column_dimensions["A"].width = 120
    wb.save(D / "instore_feedback_survey_export_20261022.xlsx")
    return {f"fb{i}": r[9] for i, r in enumerate(FR.ROWS) if len(r[9].split()) > 8}


def build_ocr():
    for fn, txt in OC.FILES.items():
        (D / fn).write_text(txt, encoding="utf-8")
    return {fn: txt for fn, txt in OC.FILES.items()}


if __name__ == "__main__":
    t = {}
    t.update(build_clienteling())
    t.update(build_feedback())
    t.update(build_ocr())
    check_texts(t)
    for p in sorted(D.iterdir()):
        print(p.name, p.stat().st_size)
