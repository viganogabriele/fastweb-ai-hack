"""python-pptx helpers (each deck defines its own look: fonts / colours / footer)."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class Deck:
    def __init__(self, theme, w=13.333, h=7.5):
        self.p = Presentation()
        self.p.slide_width = Inches(w)
        self.p.slide_height = Inches(h)
        self.W, self.H = w, h
        self.t = theme  # dict: font, title_font, fg, accent, muted, bg, footer
        self.n = 0

    def slide(self, bg=None):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        self.n += 1
        if bg:
            f = s.background.fill
            f.solid()
            f.fore_color.rgb = rgb(bg)
        return s

    def text(self, s, x, y, w, h, text, size=14, bold=False, color=None, font=None, align=None, italic=False,
             anchor=None, bullets=False, spacing=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if anchor:
            tf.vertical_anchor = anchor
        items = text if isinstance(text, list) else [text]
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            runs = it if isinstance(it, list) else [(it, {})] if isinstance(it, str) else [it]
            for rt in runs:
                if isinstance(rt, str):
                    rt = (rt, {})
                r = p.add_run()
                r.text = (("• " if bullets and rt is runs[0] else "") + rt[0])
                o = rt[1]
                r.font.size = Pt(o.get("size", size))
                r.font.bold = o.get("bold", bold)
                r.font.italic = o.get("italic", italic)
                r.font.name = o.get("font", font or self.t["font"])
                c = o.get("color", color or self.t["fg"])
                r.font.color.rgb = rgb(c)
            if align:
                p.alignment = align
            if spacing:
                p.space_after = Pt(spacing)
        return tb

    def rect(self, s, x, y, w, h, fill, line=None):
        from pptx.enum.shapes import MSO_SHAPE
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(fill)
        if line:
            r.line.color.rgb = rgb(line)
        else:
            r.line.fill.background()
        r.shadow.inherit = False
        return r

    def table(self, s, x, y, w, rows, col_w=None, size=10, header_fill=None, header_fg="FFFFFF", row_h=0.28,
              zebra=None, bold_last=False, font=None):
        nr, nc = len(rows), len(rows[0])
        gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
        tbl = gt.table
        if col_w:
            for i, cw in enumerate(col_w):
                tbl.columns[i].width = Inches(cw)
        for i in range(nr):
            tbl.rows[i].height = Inches(row_h)
            for j in range(nc):
                c = tbl.cell(i, j)
                c.margin_left = c.margin_right = Inches(0.05)
                c.margin_top = c.margin_bottom = Inches(0.02)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = c.text_frame
                tf.text = ""
                r = tf.paragraphs[0].add_run()
                r.text = str(rows[i][j])
                r.font.size = Pt(size)
                r.font.name = font or self.t["font"]
                r.font.bold = (i == 0) or (bold_last and i == nr - 1)
                r.font.color.rgb = rgb(header_fg if i == 0 else self.t["fg"])
                txt_ = str(rows[i][j]).strip()
                isnum = bool(txt_) and (txt_[0] in '+-0123456789.(' or txt_.endswith('%')) and j > 0
                tf.paragraphs[0].alignment = PP_ALIGN.RIGHT if isnum or (i == 0 and j > 0 and all(str(rr[j]).strip()[:1] in '+-0123456789.(' for rr in rows[1:] if str(rr[j]).strip())) else PP_ALIGN.LEFT
                c.fill.solid()
                if i == 0:
                    c.fill.fore_color.rgb = rgb(header_fill or self.t["accent"])
                elif zebra and i % 2 == 0:
                    c.fill.fore_color.rgb = rgb(zebra)
                else:
                    c.fill.fore_color.rgb = rgb("FFFFFF")
        return tbl

    def chart(self, s, x, y, w, h, kind, cats, series, colors=None, title=None, legend=True, labels=False,
              number_format='#,##0', font_size=10, gap=80):
        cd = CategoryChartData()
        cd.categories = cats
        for name, vals in series:
            cd.add_series(name, vals)
        gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.font.size = Pt(font_size)
        ch.font.name = self.t["font"]
        ch.has_legend = legend and len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = title
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(font_size + 1)
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
        else:
            ch.has_title = False
        plot = ch.plots[0]
        if hasattr(plot, "gap_width"):
            plot.gap_width = gap
        if labels:
            plot.has_data_labels = True
            plot.data_labels.number_format = number_format
            plot.data_labels.number_format_is_linked = False
            plot.data_labels.font.size = Pt(font_size - 1)
        if colors:
            for i, ser in enumerate(plot.series):
                c = colors[i % len(colors)]
                if kind in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
                    ser.format.line.color.rgb = rgb(c)
                    ser.smooth = False
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = rgb(c)
        try:
            ch.value_axis.has_major_gridlines = True
            ch.value_axis.major_gridlines.format.line.color.rgb = rgb("E0E0E0")
            ch.value_axis.tick_labels.font.size = Pt(font_size - 1)
            ch.category_axis.tick_labels.font.size = Pt(font_size - 1)
            ch.value_axis.tick_labels.number_format = number_format
            ch.value_axis.tick_labels.number_format_is_linked = False
        except Exception:
            pass
        return ch

    def notes(self, s, text):
        s.notes_slide.notes_text_frame.text = text

    def save(self, path):
        self.p.save(path)


__all__ = ["Deck", "rgb", "XL_CHART_TYPE", "PP_ALIGN", "MSO_ANCHOR", "Inches", "Pt", "XL_LABEL_POSITION"]
