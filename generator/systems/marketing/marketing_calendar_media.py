"""Marketing: FW26 campaign calendar (XLSX gantt), media plan by market (XLSX), agency flash report (PDF)."""
import datetime as dt
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import common as C  # noqa: E402
import config  # noqa: E402

R = C.rng("marketing_cal")
OUT = C.OUT / "marketing"
OUT.mkdir(parents=True, exist_ok=True)

# activity: id, market group, activity, type, channel, start, end, owner, budget_eur_k, status, note
ACT = [
    ("MK-0101", "US", "FW26 'Altitude' brand campaign - OOH NYC/LA", "CAMPAIGN", "OOH", "2026-09-15", "2026-10-31", "L. Mitchell / Hartwell & Rowe", 1450, "LIVE", "Times Sq + Soho wallscapes, Rodeo Dr billboard"),
    ("MK-0102", "US", "FW26 'Altitude' - digital (Meta/IG/YouTube/CTV)", "CAMPAIGN", "Digital", "2026-09-15", "2026-11-08", "Hartwell & Rowe", 980, "LIVE", ""),
    ("MK-0103", "UK", "FW26 'Altitude' - London Underground + Harrods windows", "CAMPAIGN", "OOH", "2026-09-15", "2026-10-25", "F. McLeod / Hartwell & Rowe", 720, "LIVE", "Harrods windows Sep 15-Oct 4"),
    ("MK-0104", "UK", "FW26 'Altitude' - digital", "CAMPAIGN", "Digital", "2026-09-15", "2026-11-08", "Hartwell & Rowe", 410, "LIVE", ""),
    ("MK-0105", "EU (IT/FR/DE/ES)", "FW26 'Altitude' - digital + print", "CAMPAIGN", "Digital/Print", "2026-10-01", "2026-11-15", "Mediaplus Milano", 1250, "LIVE", "Vogue IT/FR Oct issues"),
    ("MK-0106", "APAC (KR/JP)", "FW26 'Altitude' - digital (Naver, Kakao, LINE, IG)", "CAMPAIGN", "Digital", "2026-10-08", "2026-11-22", "Pacific Spark", 690, "LIVE", ""),
    ("MK-0107", "China", "FW26 'Altitude' - WeChat / Xiaohongshu / Weibo", "CAMPAIGN", "Social", "2026-10-08", "2026-11-11", "Pacific Spark SH", 880, "LIVE", ""),
    ("MK-0108", "Australia", "FW26 'Altitude' - winter campaign (global assets)", "CAMPAIGN", "Digital", "2026-09-15", "2026-10-31", "D. Lim", 85, "LIVE", "global winter assets (down jackets, snow)"),
    ("MK-0201", "Global", "Valentino Couture ASTRA capsule launch", "PRODUCT_LAUNCH", "All", "2026-09-24", "2026-10-25", "Couture team", 1100, "LIVE", "drop Thursday 24/09 10:00 local"),
    ("MK-0202", "China", "ASTRA VIP preview - Shanghai Plaza 66", "CLIENT_EVENT", "Event", "2026-09-26", "2026-09-26", "Zhao Lei", 95, "DONE", ""),
    ("MK-0203", "HK", "ASTRA window takeover Canton Road", "VM", "Window", "2026-09-24", "2026-10-21", "Vivian Lau", 40, "LIVE", ""),
    ("MK-0204", "Global", "FW26 collection launch (IRIDIA, AMALIA, ARTEM, RIVIERA)", "PRODUCT_LAUNCH", "All", "2026-09-03", "2026-09-30", "Merch / Marketing", 300, "DONE", ""),
    ("MK-0301", "Italy", "Milano Fashion Week - MIL01 VIC appointments + window", "CLIENT_EVENT", "Store", "2026-09-22", "2026-09-28", "Giulia Ferrandi", 60, "DONE", ""),
    ("MK-0302", "France", "Paris Fashion Week - PAR01 private viewing", "CLIENT_EVENT", "Store", "2026-09-29", "2026-10-07", "Camille Lefèvre", 75, "DONE", "private viewing Sat 03/10"),
    ("MK-0303", "UK", "LON01 Private Client Evening - FW26 preview", "CLIENT_EVENT", "Store", "2026-10-08", "2026-10-08", "Olivia Hartley / CRM UK", 48, "DONE", "45 VIC invited, 18:30-21:30"),
    ("MK-0304", "Canada", "TIFF gifting lounge (TOR01)", "PR", "Event", "2026-09-10", "2026-09-20", "Émilie Gagnon", 55, "DONE", ""),
    ("MK-0305", "Korea", "BIFF Busan - BUS01 pop-up corner", "CLIENT_EVENT", "Store", "2026-10-01", "2026-10-10", "Yoon Sung-ho", 38, "DONE", ""),
    ("MK-0306", "Singapore", "F1 Singapore GP weekend - SIN01 client cocktail", "CLIENT_EVENT", "Store", "2026-10-09", "2026-10-11", "Grace Tan", 30, "DONE", ""),
    ("MK-0401", "China outbound", "Golden Week clienteling (WeChat) - tax-free Milano/Paris/London", "CRM", "WeChat", "2026-09-28", "2026-10-08", "CRM Global", 25, "DONE", ""),
    ("MK-0402", "Japan", "Price increase client communication (DM + LINE) - effective 01/11", "CRM", "DM/LINE", "2026-10-09", "2026-10-31", "Takuya Suzuki", 12, "LIVE", "letter dated 09/10"),
    ("MK-0403", "China", "Double 11 - Tmall presale", "ECOM", "Tmall", "2026-10-20", "2026-11-11", "E-com China", 260, "LIVE", "stores NOT price matching"),
    ("MK-0404", "Global", "Holiday campaign 'Gift of Warmth'", "CAMPAIGN", "All", "2026-11-16", "2026-12-31", "Global Marketing", 2100, "PLANNED", ""),
    ("MK-0405", "Global", "Atelier ski season opening", "CAMPAIGN", "Digital/Store", "2026-11-23", "2026-12-20", "Global Marketing", 420, "PLANNED", ""),
    ("MK-0406", "Korea/Japan/China", "IRIDIA 'Pink edit' social push", "SOCIAL", "IG/Kakao/XHS", "2026-10-06", "2026-10-25", "C. Valli", 35, "LIVE", "added 06/10 after K-pop moment - NO STOCK until 20/11!"),
    ("MK-0407", "Global", "CRM newsletter - FW26 icons (Nova, Aurea, Liviana)", "CRM", "Email", "2026-09-17", "2026-09-17", "CRM Global", 8, "DONE", "Aurea black featured as hero (DC out!)"),
    ("MK-0408", "Global", "CRM newsletter - Couture Astra", "CRM", "Email", "2026-09-24", "2026-09-24", "CRM Global", 8, "DONE", ""),
    ("MK-0409", "Global", "CRM newsletter - Cold weather edit", "CRM", "Email", "2026-10-14", "2026-10-14", "CRM Global", 8, "DONE", "sent EU only"),
]
TYPE_FILL = {"CAMPAIGN": "1F3864", "PRODUCT_LAUNCH": "C00000", "CLIENT_EVENT": "BF8F00", "CRM": "548235", "PR": "7030A0",
             "VM": "808080", "ECOM": "ED7D31", "SOCIAL": "FF66CC"}

wb = Workbook()
ws = wb.active
ws.title = "Calendar FW26"
weeks = [dt.date(2026, 7, 27) + dt.timedelta(weeks=i) for i in range(23)]   # W31..W53
FIRSTC = 7
ws["A1"] = "VALENTINO FW26 - MARKETING CALENDAR"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "v7 - aggiornato 06/10/2026 (C. Valli). Colori: blu campagna, rosso lancio prodotto, oro eventi clienti, verde CRM, viola PR, arancio e-com, rosa social"
heads = ["ID", "Mercato", "Attività", "Tipo", "Inizio", "Fine"]
for j, h in enumerate(heads):
    c = ws.cell(row=4, column=j + 1, value=h)
    c.font = Font(bold=True)
    ws.merge_cells(start_row=4, start_column=j + 1, end_row=5, end_column=j + 1)
# month header merged
month_cols = {}
for i, w in enumerate(weeks):
    month_cols.setdefault(w.strftime("%b %Y"), []).append(FIRSTC + i)
    c = ws.cell(row=5, column=FIRSTC + i, value=f"W{w.isocalendar()[1]}\n{w:%d/%m}")
    c.alignment = Alignment(wrap_text=True, horizontal="center")
    c.font = Font(size=8)
    ws.column_dimensions[get_column_letter(FIRSTC + i)].width = 6.5
for mth, cols in month_cols.items():
    ws.cell(row=4, column=cols[0], value=mth.upper()).alignment = Alignment(horizontal="center")
    ws.cell(row=4, column=cols[0]).font = Font(bold=True, color="FFFFFF")
    ws.cell(row=4, column=cols[0]).fill = PatternFill("solid", fgColor="404040")
    if len(cols) > 1:
        ws.merge_cells(start_row=4, start_column=cols[0], end_row=4, end_column=cols[-1])
row = 6
by_market = {}
for a in ACT:
    by_market.setdefault(a[1], []).append(a)
for mkt, acts in by_market.items():
    ws.cell(row=row, column=1, value=mkt.upper()).font = Font(bold=True)
    ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor="D9D9D9")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    row += 1
    for a in acts:
        ws.cell(row=row, column=1, value=a[0])
        ws.cell(row=row, column=2, value=a[1])
        ws.cell(row=row, column=3, value=a[2])
        ws.cell(row=row, column=4, value=a[3])
        ws.cell(row=row, column=5, value=dt.date.fromisoformat(a[5])).number_format = "DD/MM/YY"
        ws.cell(row=row, column=6, value=dt.date.fromisoformat(a[6])).number_format = "DD/MM/YY"
        s, e = dt.date.fromisoformat(a[5]), dt.date.fromisoformat(a[6])
        for i, w in enumerate(weeks):
            if w <= e and w + dt.timedelta(days=6) >= s:
                c = ws.cell(row=row, column=FIRSTC + i)
                c.fill = PatternFill("solid", fgColor=TYPE_FILL[a[3]])
        if a[10]:
            ws.cell(row=row, column=3).comment = Comment(a[10], a[7].split("/")[0].strip())
        row += 1
ws.column_dimensions["A"].width = 9
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 58
ws.column_dimensions["D"].width = 14
ws.freeze_panes = "G6"
ws.cell(row=row + 1, column=3, value="Golden Week CN 01-08/10, Chuseok KR 24-26/09, Double 11 11/11, Japan price +8% 01/11")
ws.cell(row=row + 1, column=3).font = Font(italic=True)

ws2 = wb.create_sheet("List")
ws2.append(["ID", "Market", "Activity", "Type", "Channel", "Start", "End", "Owner", "Budget EUR k", "Status", "Notes"])
for a in ACT:
    ws2.append([a[0], a[1], a[2], a[3], a[4], a[5], a[6], a[7], a[8], a[9], a[10]])
ws2.append(["", "", "TOTAL", "", "", "", "", "", f"=SUM(I2:I{len(ACT) + 1})", "", ""])
for c in ws2[1]:
    c.font = Font(bold=True)
ws2.column_dimensions["C"].width = 60
ws2.column_dimensions["K"].width = 50

ws3 = wb.create_sheet("v3_OLD")
ws3["A1"] = "v3 - 30/07 - SUPERSEDED"
old = [("MK-0101", "US OOH", "2026-09-22", "2026-10-31"), ("MK-0103", "UK OOH", "2026-09-22", "2026-10-25"),
       ("MK-0201", "Couture ASTRA launch", "2026-10-01", "2026-10-31"), ("MK-0303", "LON01 client evening", "2026-10-15", "2026-10-15"),
       ("MK-0402", "Japan price increase comms", "2026-10-15", "2026-10-31"), ("MK-0106", "APAC digital", "2026-09-28", "2026-11-15")]
ws3.append(["ID", "Activity", "Start", "End"])
for o in old:
    ws3.append(list(o))
wb.save(OUT / "campaign_calendar_FW26.xlsx")

# ------------------------------------------------------------------ media plan by market
MP = []
markets = {
    "US": ("USD", "Hartwell & Rowe Media", [("OOH", 1100), ("Meta/IG", 420), ("YouTube", 230), ("CTV", 210), ("Search", 120), ("Print", 160)]),
    "UK": ("GBP", "Hartwell & Rowe Media", [("OOH", 520), ("Meta/IG", 150), ("YouTube", 70), ("Search", 60), ("Print", 90)]),
    "IT": ("EUR", "Mediaplus Milano", [("Print", 180), ("Meta/IG", 140), ("Programmatic", 90), ("Search", 50)]),
    "FR": ("EUR", "Mediaplus Milano", [("Print", 170), ("Meta/IG", 130), ("Programmatic", 80), ("Search", 45)]),
    "DE": ("EUR", "Mediaplus Milano", [("Meta/IG", 110), ("Programmatic", 70), ("Search", 40)]),
    "ES": ("EUR", "Mediaplus Milano", [("Meta/IG", 60), ("Search", 25)]),
    "KR": ("KRW", "Pacific Spark", [("Naver", 160), ("Kakao", 110), ("Meta/IG", 120), ("YouTube", 60)]),
    "JP": ("JPY", "Pacific Spark", [("LINE", 90), ("Meta/IG", 110), ("Print", 80), ("Yahoo! JAPAN", 40)]),
    "CN": ("CNY", "Pacific Spark SH", [("WeChat", 300), ("Xiaohongshu", 260), ("Weibo", 120), ("Douyin", 200)]),
    "AU": ("AUD", "in-house", [("Meta/IG", 55), ("Search", 30)]),
}
START = {"US": "2026-09-15", "UK": "2026-09-15", "IT": "2026-10-01", "FR": "2026-10-01", "DE": "2026-10-01", "ES": "2026-10-01",
         "KR": "2026-10-08", "JP": "2026-10-08", "CN": "2026-10-08", "AU": "2026-09-15"}
for m, (ccy, agency, chans) in markets.items():
    for ch, eur_k in chans:
        cpm = {"OOH": 9, "Print": 28, "CTV": 32, "YouTube": 14, "Search": 0, "Meta/IG": 9.5, "Programmatic": 6, "Naver": 11, "Kakao": 10,
               "LINE": 8, "Yahoo! JAPAN": 7, "WeChat": 12, "Xiaohongshu": 15, "Weibo": 6, "Douyin": 10}[ch]
        imp = int(eur_k * 1000 / cpm * 1000) if cpm else ""
        delivered = round(R.uniform(0.55, 0.85), 2) if START[m] <= "2026-09-15" else round(R.uniform(0.2, 0.5), 2)
        MP.append(dict(market=m, agency=agency, channel=ch, flight_start=START[m],
                       flight_end="2026-11-08" if ch not in ("OOH", "Print") else "2026-10-31",
                       budget_local=round(eur_k * 1000 * config.FX_EUR[ccy], -2), currency=ccy, budget_eur=eur_k * 1000,
                       planned_impressions=imp, planned_cpm_eur=cpm or "", kpi="Reach/Frequency" if ch in ("OOH", "Print", "CTV") else (
                           "CPC / ROAS" if ch == "Search" else "VTR / engagement"),
                       spend_to_date_pct=delivered, notes=""))
mp = pd.DataFrame(MP)
mp.loc[(mp.market == "AU"), "notes"] = "same creative as EU/US (snow, down) - spring in AU"
mp.loc[(mp.market == "UK") & (mp.channel == "OOH"), "notes"] = "incl. Harrods windows (barter, not in budget)"
with pd.ExcelWriter(OUT / "media_plan_FW26_by_market.xlsx", engine="openpyxl") as xw:
    mp.to_excel(xw, sheet_name="MediaPlan", index=False)
    mp.groupby(["market", "currency"]).agg(budget_local=("budget_local", "sum"), budget_eur=("budget_eur", "sum")).reset_index().to_excel(
        xw, sheet_name="ByMarket", index=False)

# ------------------------------------------------------------------ agency flash report (PDF)
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

k = pd.read_csv(C.INTERNAL / "kpi" / "store_daily_kpi.csv")
camp = ["NYC01", "NYC02", "BHS01", "MIA01", "LON01", "LON02"]


def tot(ss, a, b):
    x = k[(k.store_id.isin(ss)) & (k.business_date >= a) & (k.business_date <= b)]
    return x.net_sales_eur_bud.sum()


rows = [["Market", "Doors", "Sales TY (EUR k)", "Sales LY (EUR k)", "YoY", "Attributed uplift*"]]
for nm, ss, att in (("US", camp[:4], "+19%"), ("UK", camp[4:], "+26%")):
    a, b = tot(ss, "2026-09-15", "2026-10-11"), tot(ss, "2025-09-15", "2025-10-11")
    rows.append([nm, ", ".join(ss), f"{a / 1000:,.0f}", f"{b / 1000:,.0f}", f"{a / b - 1:+.1%}", att])
a, b = tot(camp, "2026-09-15", "2026-10-11"), tot(camp, "2025-09-15", "2025-10-11")
rows.append(["TOTAL", "6 doors + brand.com US/UK", f"{a / 1000:,.0f}", f"{b / 1000:,.0f}", f"{a / b - 1:+.1%}", "+22%"])

ss_ = getSampleStyleSheet()
doc = SimpleDocTemplate(str(OUT / "HartwellRowe_FW26_Altitude_flash_results_US_UK_wk38-41.pdf"), pagesize=A4,
                        title="FW26 Altitude - Flash Results US & UK", author="Hartwell & Rowe Media")
story = [Paragraph("HARTWELL &amp; ROWE MEDIA", ss_["Title"]),
         Paragraph("Valentino FW26 'Altitude' - Flash Results, US &amp; UK (weeks 38-41, 15 Sep - 11 Oct 2026)", ss_["Heading2"]),
         Spacer(1, 8),
         Paragraph("<b>Headline: the campaign delivered a +22% attributed sales uplift in the US and UK.</b> "
                   "OOH reach exceeded plan (NYC 78% of A25-54 HHI$250k+, London 71%), digital VTR 41% vs 30% benchmark.", ss_["BodyText"]),
         Spacer(1, 10)]
t = Table(rows, colWidths=[55, 150, 80, 80, 50, 80])
t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                       ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
story += [t, Spacer(1, 10),
          Paragraph("Key drivers", ss_["Heading3"]),
          Paragraph("- London: exceptional week 41 driven by brand heat (New Bond St +3x on 8 Oct).<br/>"
                    "- New York: Madison Ave resilient despite UNGA week; SoHo strong on hero icons.<br/>"
                    "- brand.com: paid social last-click revenue +31% YoY.", ss_["BodyText"]),
          Spacer(1, 10),
          Paragraph("Methodology", ss_["Heading3"]),
          Paragraph("*Attributed uplift = (TY sales in exposed doors + last-click attributed e-commerce revenue) vs LY same weeks, "
                    "adjusted by the H&amp;R Luxury Halo Index (1.35). No holdout markets were defined for this flight. "
                    "Store sales provided by client in EUR at budget rate; LY base includes the 2025 period with no campaign activity. "
                    "Client events and CRM activations during the flight are included in exposed-door sales.", ss_["BodyText"]),
          Spacer(1, 6),
          Paragraph("Confidential - prepared for Valentino Marketing US/UK. 14 Oct 2026.", ss_["Italic"])]
doc.build(story)
print("marketing calendar/media/report done")
