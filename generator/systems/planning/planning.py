"""Merchandise planning sources: OTB workbook, new-product allocation plan, replenishment parameters."""
import datetime as dt
import glob
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import common as C  # noqa: E402
import config  # noqa: E402

R = C.rng("planning")
OUT = C.OUT / "planning"
OUT.mkdir(parents=True, exist_ok=True)
st = C.stores().set_index("store_id")
sk = C.skus().set_index("sku")

# ------------------------------------------------------------------ actuals by region x category x month (EUR k, budget rate)
fs = sorted(glob.glob(str(C.INTERNAL / "sales" / "daily_store_sku" / "*.csv")))
fs = [f for f in fs if any(y in f for y in ("2025_08", "2025_09", "2025_10", "2025_11", "2025_12", "2026_01", "2026_08", "2026_09", "2026_10"))]
fs += sorted(glob.glob(str(C.INTERNAL / "sales" / "daily_store_sku" / "sales_daily_2025_0[1]*.csv")))
sales = pd.concat([pd.read_csv(f) for f in fs])
sales = sales.merge(C.stores()[["store_id", "region"]]).merge(C.skus()[["sku", "category"]])
sales["eur_k"] = sales.net_sales_lc / sales.currency.map(config.FX_EUR) / 1000
sales["m"] = sales.business_date.str[:7]
act = sales.groupby(["region", "category", "m"]).eur_k.sum()

MONTHS = ["2026-08", "2026-09", "2026-10", "2026-11", "2026-12", "2027-01"]
LY = {m: f"{int(m[:4]) - 1}{m[4:]}" for m in MONTHS}
CATS = ["Outerwear", "Knitwear", "Jersey", "Footwear", "Accessories", "Bottoms"]
REGIONS = ["EMEA", "AMERICAS", "APAC"]
GROWTH = {"EMEA": 1.05, "AMERICAS": 1.04, "APAC": 1.09}
CAT_ADJ = {"Outerwear": 1.03, "Knitwear": 1.02, "Jersey": 0.98, "Footwear": 1.0, "Accessories": 1.04, "Bottoms": 1.0}
WOC = {"Outerwear": 2.4, "Knitwear": 2.0, "Jersey": 1.6, "Footwear": 2.2, "Accessories": 1.8, "Bottoms": 2.5}   # months of cover EOP


def ly(reg, cat, m):
    try:
        return float(act.loc[(reg, cat, LY[m])])
    except KeyError:
        return 0.0


def actual(reg, cat, m):
    try:
        return float(act.loc[(reg, cat, m)])
    except KeyError:
        return None


# ------------------------------------------------------------------ OTB workbook
wb = Workbook()
thin = Side(style="thin", color="999999")
HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
ACT = PatternFill("solid", fgColor="E2EFDA")
OTBF = PatternFill("solid", fgColor="FCE4D6")
WARN = PatternFill("solid", fgColor="FFF2CC")
KFMT = '#,##0;[Red]-#,##0'

ws0 = wb.active
ws0.title = "Assumptions"
ws0["A1"] = "FW26 Merchandise Plan / OTB - Assunzioni"
ws0["A1"].font = Font(bold=True, size=14)
ws0["A2"] = "Valori in EUR k al cambio budget 2026, retail value. Versione v5 - 14/09/2026 (G. Bellini)"
ws0.append([])
ws0.append(["Regione", "Crescita vs LY", "Note"])
for r in REGIONS:
    ws0.append([r, GROWTH[r], {"EMEA": "incl. Golden Week Milano/Parigi", "AMERICAS": "campagna FW26 US/UK da 15/9 (+ agenzia stima +20%?)",
                               "APAC": "Genius Polaris + Cina forte"}[r]])
ws0.append([])
ws0.append(["Categoria", "Adj. crescita", "Mesi copertura EOP (WOC/4)"])
for c in CATS:
    ws0.append([c, CAT_ADJ[c], WOC[c]])
ws0.append([])
ws0.append(["Polaris (Genius)", "Piano EU = profilo Genius drop 2025 (CN/HK), fattore 1.0", "da rivedere dopo lancio 24/9"])
ws0.append(["ILVARO", "analogo BADY x 0,55 (allocazione) / forecast stat. BADY x 0,6", "colore 510 rosa cipria = 20% buy"])
ws0.append(["SYD01", "profilo stagionalità NH_STD (come EMEA)", "?? chiedere a Daniel Lim"])
for row in ws0.iter_rows(min_row=4, max_row=ws0.max_row):
    for c in row:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
ws0.column_dimensions["A"].width = 22
ws0.column_dimensions["B"].width = 55
ws0.column_dimensions["C"].width = 45
ws0["C14"].fill = WARN

MEASURES = ["Planned Sales", "Actual Sales", "Var % vs plan", "BOP", "Planned EOP", "Planned Receipts", "On Order", "OTB"]
sheet_rows = {}
for reg in REGIONS:
    ws = wb.create_sheet(reg)
    ws["A1"] = f"OTB FW26 - {reg}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "EUR k - retail @ budget FX"
    # header: months merged by quarter
    ws.cell(row=3, column=3, value="Q3 FY26").fill = HDR
    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column=4)
    ws.cell(row=3, column=5, value="Q4 FY26").fill = HDR
    ws.merge_cells(start_row=3, start_column=5, end_row=3, end_column=7)
    ws.cell(row=3, column=8, value="Q1 FY27").fill = HDR
    ws.cell(row=3, column=9, value="Season").fill = HDR
    for col in range(3, 10):
        ws.cell(row=3, column=col).font = Font(color="FFFFFF", bold=True)
        ws.cell(row=3, column=col).alignment = Alignment(horizontal="center")
    labels = ["Ago-26", "Set-26", "Ott-26", "Nov-26", "Dic-26", "Gen-27", "TOT"]
    ws.cell(row=4, column=1, value="Categoria")
    ws.cell(row=4, column=2, value="Misura")
    for j, l in enumerate(labels):
        c = ws.cell(row=4, column=3 + j, value=l)
        c.font = Font(bold=True)
        c.fill = SUB
    r = 5
    blocks = {}
    for cat in CATS:
        blocks[cat] = r
        ws.cell(row=r, column=1, value=cat).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r + len(MEASURES) - 1, end_column=1)
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center")
        plan = [round(ly(reg, cat, m) * GROWTH[reg] * CAT_ADJ[cat] * R.uniform(0.97, 1.03)) for m in MONTHS]
        nxt = plan[1:] + [round(plan[-1] * 0.7)]
        eop = [round(n * WOC[cat]) for n in nxt]
        bop0 = round(plan[0] * WOC[cat] * R.uniform(0.85, 1.05))
        for k, meas in enumerate(MEASURES):
            rr = r + k
            ws.cell(row=rr, column=2, value=meas)
            for j, m in enumerate(MONTHS):
                col = 3 + j
                L = get_column_letter(col)
                cell = ws.cell(row=rr, column=col)
                if meas == "Planned Sales":
                    cell.value = plan[j]
                elif meas == "Actual Sales":
                    a = actual(reg, cat, m)
                    if m <= "2026-10" and a is not None:
                        cell.value = round(a)
                        cell.fill = ACT
                        if m == "2026-10":
                            cell.comment = Comment("MTD al 21/10", "planning")
                elif meas == "Var % vs plan":
                    cell.value = f'=IF({L}{r + 1}="","",{L}{r + 1}/{L}{r}-1)'
                    cell.number_format = "0.0%"
                elif meas == "BOP":
                    cell.value = bop0 if j == 0 else f"={get_column_letter(col - 1)}{r + 4}"
                elif meas == "Planned EOP":
                    cell.value = eop[j]
                elif meas == "Planned Receipts":
                    cell.value = f"={L}{r}+{L}{r + 4}-{L}{r + 3}"
                elif meas == "On Order":
                    rec_est = plan[j] + eop[j] - (bop0 if j == 0 else eop[j - 1])
                    frac = [1.0, 1.0, 0.92, 0.7, 0.45, 0.2][j]
                    cell.value = max(0, round(rec_est * frac * R.uniform(0.95, 1.05)))
                elif meas == "OTB":
                    cell.value = f"={L}{r + 5}-{L}{r + 6}"
                    cell.fill = OTBF
                if meas not in ("Var % vs plan",):
                    cell.number_format = KFMT
            # season total
            tcell = ws.cell(row=rr, column=9)
            if meas in ("Planned Sales", "Actual Sales", "Planned Receipts", "On Order", "OTB"):
                tcell.value = f"=SUM(C{rr}:H{rr})"
                tcell.number_format = KFMT
            elif meas == "BOP":
                tcell.value = f"=C{rr}"
                tcell.number_format = KFMT
            elif meas == "Planned EOP":
                tcell.value = f"=H{rr}"
                tcell.number_format = KFMT
            elif meas == "Var % vs plan":
                tcell.value = f'=IF(SUM(C{r + 1}:E{r + 1})=0,"",SUM(C{r + 1}:E{r + 1})/SUM(C{r}:E{r})-1)'
                tcell.number_format = "0.0%"
        r += len(MEASURES)
    # region total with SUM over categories
    tot = r
    ws.cell(row=tot, column=1, value="TOTALE " + reg).font = Font(bold=True)
    ws.merge_cells(start_row=tot, start_column=1, end_row=tot + len(MEASURES) - 1, end_column=1)
    for k, meas in enumerate(MEASURES):
        rr = tot + k
        ws.cell(row=rr, column=2, value=meas).font = Font(bold=True)
        for col in range(3, 10):
            L = get_column_letter(col)
            if meas == "Var % vs plan":
                ws.cell(row=rr, column=col, value=f'=IF({L}{tot + 1}=0,"",{L}{tot + 1}/{L}{tot}-1)').number_format = "0.0%"
            else:
                refs = ",".join(f"{L}{blocks[c] + k}" for c in CATS)
                ws.cell(row=rr, column=col, value=f"=SUM({refs})").number_format = KFMT
            ws.cell(row=rr, column=col).font = Font(bold=True)
    sheet_rows[reg] = (blocks, tot)
    ws.freeze_panes = "C5"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 18
    for col in range(3, 10):
        ws.column_dimensions[get_column_letter(col)].width = 11
    # planner comments
    if reg == "EMEA":
        c = ws.cell(row=blocks["Outerwear"] + 6, column=5)
        c.comment = Comment("Include riassortimento MAIRE nero 420 pz (PO 4500218734) - ETA 02/10 (era 25/9)", "G. Bellini")
    if reg == "APAC":
        c = ws.cell(row=blocks["Outerwear"], column=4)
        c.comment = Comment("Polaris: piano CN/HK ok. ILVARO rosa nessun upside previsto", "W. Chen")

# Summary sheet with cross-sheet formulas
ws = wb.create_sheet("Summary", 0)
ws["A1"] = "FW26 OTB - Summary (EUR k)"
ws["A1"].font = Font(bold=True, size=14)
ws.append([])
ws.append(["Regione", "Planned Sales", "Actual Sales (Aug-Oct MTD)", "Plan Aug-Oct", "Var %", "Planned Receipts", "On Order", "OTB residuo"])
for c in ws[3]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HDR
for i, reg in enumerate(REGIONS):
    blocks, tot = sheet_rows[reg]
    rr = 4 + i
    ws.append([reg, f"='{reg}'!I{tot}", f"='{reg}'!I{tot + 1}", f"=SUM('{reg}'!C{tot}:E{tot})", f"=C{rr}/D{rr}-1",
               f"='{reg}'!I{tot + 5}", f"='{reg}'!I{tot + 6}", f"='{reg}'!I{tot + 7}"])
ws.append(["TOTALE", "=SUM(B4:B6)", "=SUM(C4:C6)", "=SUM(D4:D6)", "=C7/D7-1", "=SUM(F4:F6)", "=SUM(G4:G6)", "=SUM(H4:H6)"])
for row in ws.iter_rows(min_row=4, max_row=7):
    for c in row[1:]:
        c.number_format = KFMT
    row[4].number_format = "0.0%"
for col in "ABCDEFGH":
    ws.column_dimensions[col].width = 20
ws["A9"] = "NB: Oct actual = MTD 21/10 vs full-month plan -> var% Oct non confrontabile"
ws["A9"].font = Font(italic=True, color="C00000")

# outdated sheet with pasted values (v3 July), hidden
ws = wb.create_sheet("OTB_v3_luglio")
ws["A1"] = "VERSIONE LUGLIO - NON USARE (valori incollati)"
ws["A1"].font = Font(bold=True, color="C00000")
ws.append(["Regione", "Categoria"] + ["Ago-26", "Set-26", "Ott-26", "Nov-26", "Dic-26", "Gen-27"])
for reg in REGIONS:
    for cat in CATS:
        ws.append([reg, cat] + [round(ly(reg, cat, m) * 1.12 * R.uniform(0.9, 1.1)) for m in MONTHS])
ws.sheet_state = "hidden"
wb.save(OUT / "FW26_merch_plan_OTB_v5.xlsx")
print("OTB saved")

# ------------------------------------------------------------------ allocation plan for new products
la = pd.read_csv(C.CACHE / "launch_allocations.csv", dtype={"size": str})
ANALOG = {"ILVARO": ("BADY", 0.55), "AMBRINE": ("HERMINE", 0.80), "ARNEB": ("MONTCLA", 0.85), "ROSIERE": ("KARAKORUM", 0.90),
          "WOOL TURTLENECK": ("CABLE KNIT SWEATER", 1.0), "POLARIS CAPE": ("GENIUS FW25 DROP (CN/HK profile)", 1.0),
          "POLARIS OVERSIZED": ("GENIUS FW25 DROP (CN/HK profile)", 1.0), "POLARIS BALACLAVA": ("GENIUS FW25 DROP (CN/HK profile)", 1.0)}
GRADE = {}
for sid in st.index:
    t = st.store_type[sid]
    GRADE[sid] = "A" if t == "Flagship" else ("C" if t == "Resort" else "B")
rows = []
for r in la.itertuples():
    name = C.STYLE_NAME[r.sku.split("-")[0]]
    an, fac = ANALOG[name]
    calc = r.qty * R.uniform(0.82, 1.18)
    run = config.SIZE_RUNS[sk.size_run[r.sku]]
    share = config.SIZE_CURVES[sk.size_run[r.sku]][run.index(r.size)]
    rows.append(dict(wave=f"FW26 launch {r.launch_date}", store=r.store_id, werks=C.SAP_WERKS[r.store_id], store_grade=GRADE[r.store_id],
                     cluster=st.subregion[r.store_id], dc=C.REGION_DC[st.region[r.store_id]], style=name, sku=r.sku, size=r.size,
                     analog_style=an, analog_factor=fac, size_curve="STD_" + sk.size_run[r.sku] + ("_ASIA" if st.region[r.store_id] == "APAC" else ""),
                     size_share=round(share, 3), weeks_cover=4, calc_qty=round(calc, 2), min_presentation=1 if GRADE[r.store_id] != "C" else 0,
                     alloc_qty=int(r.qty), override="" if abs(calc - r.qty) < 1.5 else "MANUAL", gi_date=r.gi_date,
                     season_profile="NH_STD"))
al = pd.DataFrame(rows)
# quirk: duplicated row (ILVARO 510 SEL02 size 1) and store names instead of codes in POLARIS sheet
dup = al[(al.sku == "J20931A00618-510") & (al.store == "SEL02") & (al["size"] == "1")]
al = pd.concat([al, dup]).sort_values(["wave", "style", "sku", "store", "size"], kind="stable")
with pd.ExcelWriter(OUT / "allocation_plan_FW26_new_products.xlsx", engine="openpyxl") as xw:
    a1 = al[al.wave.str.contains("2026-09-03")].copy()
    a1.to_excel(xw, sheet_name="Wave1_0309", index=False)
    a2 = al[al.wave.str.contains("2026-09-24")].copy()
    a2["store"] = a2.store.map(lambda s: str(R.choice(C.STORE_ALIASES[s])) if R.random() < 0.6 else st.store_name[s])
    a2 = a2.drop(columns=["werks"])
    a2.to_excel(xw, sheet_name="Wave2_Genius_2409", index=False)
    summ = al.groupby(["wave", "style", "sku", "dc"]).alloc_qty.sum().unstack("dc").fillna(0).astype(int).reset_index()
    summ.to_excel(xw, sheet_name="Summary_by_DC", index=False)
    pd.DataFrame({"parametro": ["metodo", "ILVARO", "AMBRINE", "POLARIS", "rounding", "min presentation", "nota SYD01", "nota rosa"],
                  "valore": ["analogo x fattore x settimane copertura x curva taglie", "BADY x 0,55 (prudenziale, nuovo fit boxy)",
                             "HERMINE x 0,8; split colore 70% nero / 30% marrone", "come drop Genius FW25 (profilo CN/HK applicato a tutte le regioni)",
                             "half-up, min 1 per taglie core in store A/B", "1 pz taglie con quota >= 12% (A/B)",
                             "profilo stagionale NH_STD come Europa", "510 rosa: buy 20% del totale ILVARO, rischio sell-through basso"]}
                 ).to_excel(xw, sheet_name="Parametri", index=False)
print("allocation rows", len(al))

# ------------------------------------------------------------------ replenishment parameters (SAP-like)
LEAD = {"Italy": 1, "France": 2, "Switzerland": 3, "UK": 3, "Germany": 2, "Spain": 2, "Benelux": 2, "Nordics": 3,
        "Middle East": 5, "USA": 5, "Canada": 6, "Japan": 6, "Korea": 5, "Mainland China": 8, "HK/Macau/TW": 5,
        "South-East Asia": 6, "Oceania": 7}
COVER = {"EMEA": 21, "AMERICAS": 28, "APAC": 28}
DAYN = ["MO", "DI", "MI", "DO", "FR", "SA", "SO"]
rp = []
for sid in st.index:
    reg = st.region[sid]
    t = st.store_type[sid]
    if reg == "APAC":
        od = [1, 4] if t == "Flagship" else [1]
    else:
        od = [0, 3] if t == "Flagship" else [0]
    lt = LEAD[st.subregion[sid]]
    changed, user = "2026-07-28", "PLN_GBELLI"
    status = "ACTIVE"
    note = ""
    if sid == "MAD01":
        lt = 1
        changed, user = "2024-02-12", "PLN_OLD_MIGR"
        note = "lead time da migrazione 2024"
    if sid == "COU01":
        status, note = "BLOCKED_SEASON", "chiuso 13/09-26/11"
    if sid == "STM01":
        note = "Zwischensaison dal 19/10 - bloccare riordini!"
    if sid == "CRT01":
        note = "ottobre solo ven-dom (giorno ordine lunedì invariato)"
    if sid == "ROM01":
        note = "cantiere via Condotti fino 6/11 - nessuna modifica parametri"
    rp.append(dict(WERKS=C.SAP_WERKS[sid], NAME1=st.store_name[sid], STORE_ID=sid, VKORG={"EMEA": "1000", "AMERICAS": "2000", "APAC": "3000"}[reg],
                   LIEF_WERK={"EMEA": "IT10", "AMERICAS": "US20", "APAC": "HK30"}[reg], BESTELLTAGE="/".join(DAYN[d_] for d_ in od),
                   WBZ_TAGE=lt, REICHWEITE_TAGE=COVER[reg], SAFETY_FACTOR=1.15, SERVICE_LEVEL=0.95 if t == "Flagship" else 0.92,
                   MIN_PRES_RULE="1 pz taglie con quota curva >= 12%" if t in ("Flagship", "Boutique") else "nessuna",
                   SAISON_PROFIL="NH_STD" if sid != "DXB01" else "TROP_STD", FCST_MODEL="STAT_HW_WEEKLY", STATUS=status,
                   CUTOFF_ORDINE="10:00 CET" if reg == "EMEA" else ("10:00 ET" if reg == "AMERICAS" else "10:00 HKT"),
                   AENDERUNG_AM=changed, AENDERUNG_VON=user, NOTE=note))
rpd = pd.DataFrame(rp)
# quirk: SYD01 seasonality profile NH_STD although southern hemisphere; the field exists with value SH for no store
rpd.to_csv(OUT / "replenishment_parameters_stores.csv", index=False, sep=";", decimal=",")

mp = []
for sid in st.index:
    t = st.store_type[sid]
    for cat in CATS:
        base = {"Flagship": 2, "Boutique": 1, "Shop-in-shop": 1, "Resort": 1}[t]
        mp.append(dict(WERKS=C.SAP_WERKS[sid], store=sid, category=cat, min_pres_per_core_size=base if cat in ("Outerwear", "Knitwear") else 1,
                       core_sizes={"Outerwear": "W:0,1,2 / M:2,3,4", "Knitwear": "S,M,L", "Jersey": "S,M,L", "Footwear": "37,38,39 / 42,43",
                                   "Accessories": "TU", "Bottoms": "M,L"}[cat],
                       max_facings=int(R.integers(4, 30)) if cat != "Bottoms" else int(R.integers(0, 6)),
                       valid_from="2026-08-01"))
mpd = pd.DataFrame(mp)
mpd.loc[(mpd.store == "SYD01") & (mpd.category == "Outerwear"), "max_facings"] = 34
mpd.loc[(mpd.store == "SYD01") & (mpd.category == "Jersey"), "max_facings"] = 6
mpd.to_csv(OUT / "min_presentation_by_store_category.csv", index=False)
print("params done")

# ------------------------------------------------------------------ recalc OTB with LibreOffice so that cached values exist
import shutil
import subprocess
import tempfile

tmp = Path(tempfile.mkdtemp())
src = OUT / "FW26_merch_plan_OTB_v5.xlsx"
try:
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp), str(src)], check=True,
                   capture_output=True, timeout=180)
    shutil.copy(tmp / src.name, src)
    print("OTB recalculated via soffice")
except Exception as e:  # noqa: BLE001
    print("soffice recalc skipped:", e)
