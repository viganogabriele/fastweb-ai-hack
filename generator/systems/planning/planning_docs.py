"""S&OP minutes / planner notes (authored text) -> MD and DOCX."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import common as C  # noqa: E402

OUT = C.OUT / "planning" / "sop_minutes"
OUT.mkdir(parents=True, exist_ok=True)

DOCS = {}

DOCS["2026-09-07_SOP_FW26_minutes"] = ("docx", """
# S&OP FW26 – riunione settimanale
**Data:** lunedì 07/09/2026, 10:00–11:30 CET (Teams)
**Presenti:** G. Bellini (Merch Planning, chair), P. Conti (Demand Planning), L. Marini (Purchasing), M. Re (DC Castel S. Giovanni), W. Chen (APAC planning, HK), S. Kowalski (Americas planning), E. Sartori (Retail Ops), R. Galli (AM Italy/Spain)
**Assenti:** D. Lim (SEA/Oceania) – mandato commenti via mail

## 1. Lancio FW26 (03/09)
- Lancio ok, allocazioni arrivate ovunque tranne CDU01 (1 giorno dopo, nessun impatto).
- Prima settimana: IRIDIA nero parte forte, AMALIA marrone vende meglio del previsto a Milano e Parigi. Split colore AMALIA era 70/30 nero/marrone (EMEA 120 vs 61 pz). Galli: "il marrone è IL colore quest'anno, abbiamo sbagliato lo split". Marini verifica se il fornitore (Romarta) può fare un piccolo riassortimento sulle taglie piccole.
- IRIDIA rosa cipria (510): partenza lenta, in linea con il piano prudenziale (buy 20% del totale IRIDIA).

## 2. AUREA nero – copertura DC
- M. Re: a IT10 restano circa 105 pz di AUREA 999, copertura fino a sabato 12/09 circa ai ritmi attuali. HK30 ~175 pz, US20 ~30 pz.
- PO di riassortimento 4500218734 (Atelier Nistru) confermato con consegna **25/09**. Marini: "il fornitore dice ok ma il tessuto laccato arriva dalla tessitura, li sollecito".
- Rischio: 2 settimane di buco sul nero. Decisione: nessuna azione, il gap è accettabile.

## 3. Couture ASTRA (lancio 24/09)
- Allocazione già chiusa con l'ipotesi "come drop Couture FW25" su tutte le regioni (fattore 1.0, profilo CN/HK). Kowalski chiede se ha senso per gli US con un prezzo così alto: "our clients didn't love the last Couture drop". Bellini: lasciamo così, rivediamo dopo la prima settimana.
- Spedizioni lancio: EMEA 21–22/09, APAC dal 16/09.

## 4. Golden Week / turismo
- Previsto +30/40% turisti cinesi a Milano e Parigi (nuovi voli Chengdu–Milano). Sartori chiede copertura extra su MIL01/PAR01 per icone (NOVA, AUREA). Conti: il forecast statistico non vede gli eventi, aggiungere override manuale.
- **Action:** Conti + Galli, override +25% MIL01/MIL02/PAR01/PAR02 per sett. 40–41. *(nota successiva: non caricato a sistema, vedi 21/09)*

## 5. Varie
- ROM01: comunicazione del Comune, cantiere davanti all'ingresso di Via Condotti dal 14/09 per ~8 settimane. Nessuna modifica ai parametri di riordino per ora.
- SYD01 (mail di D. Lim): "the forecast keeps sending us heavy down for spring, customers ask polos and tees". Parcheggiato, da discutere con Conti offline.
- Giappone: aumento prezzi 1/11 (comunicazione ai clienti 09/10). W. Chen: aspettarsi acquisti anticipati, controllare stock icone TYO01/TYO02/OSA01.

## Action items
| # | Azione | Owner | Scadenza |
|---|--------|-------|----------|
| 1 | Sollecito fornitore AUREA nero (tessuto) | L. Marini | 11/09 |
| 2 | Verifica riassortimento AMALIA marrone taglie 0-2 | L. Marini | 14/09 |
| 3 | Override Golden Week MIL/PAR | P. Conti / R. Galli | 14/09 |
| 4 | SYD01 forecast – call con D. Lim | P. Conti | 18/09 |
""")

DOCS["2026-09-21_SOP_FW26_minutes"] = ("md", """
# S&OP FW26 – weekly
Mon 21/09/2026 10:00 CET – Teams. Chair G. Bellini. Minutes: A. Ferri (planning intern)

Attendees: Bellini, Conti, Marini, M. Re, W. Chen, S. Kowalski, E. Sartori, I. Moreau (AM France), D. Lim (first 20 min)

**1. AUREA black – DC out of stock since 12/09**
- All three DCs at zero on 999. PO 4500218734: vendor moved delivery from 25/09 to **02/10** (mail of 10/09, dyeing lot rejected at the mill).
- Moreau + Galli (by mail): Korean customers in PAR01/PAR02 and MIL01 asking for black Aurea every day since the weekend, sizes 0-1-2 mostly. Stores already at zero or 1-2 pcs. Beige/navy not an alternative ("elles veulent le noir, point").
- Discussed inter-store transfers. Agreed: IBT only domestic (same country) to avoid customs paperwork; cross-border CH/UK not allowed. Moreau will check Cannes for PAR01.
- Bellini: when the 02/10 receipt lands, priority to PAR01, PAR02, MIL01, LON02.

**2. IRIDIA – top-up 600 pcs powder pink (PO 4500219960, lands 08/10)**
- Sell-through after 2.5 weeks: black ~70%, off-white ~50%, pink ~45% of the launch allocation. Pink is the weakest.
- Split of the 600 pcs decided: **HK30 330 / IT10 190 / US20 80** (APAC had the bigger launch share and better pink rate in SEL/TYO).
- E-commerce asks for a share of the top-up; decision postponed to when goods land.
- Chen: some talk on Korean social about a celebrity at Milan FW wearing Valentino, nothing specific on Iridia.

**3. AMALIA brown**
- Sizes 0-1-2 almost gone in EMEA stores. Marini: Romarta can do 230 pcs small sizes but not before ~10/11. To be raised this week (PO).

**4. Golden Week override**
- Not loaded in time for week 40 (Conti: "the template rejected the file, retrying"). Retail Ops not happy.

**5. Couture ASTRA** – launch Thursday 24/09. Nothing to add.

**6. SYD01** – Lim: Sydney is in spring, 22-25 degrees, the store sells tees/polos/light jackets and keeps receiving puffers because the forecast uses the northern seasonality. Conti: changing the season profile requires a model rebuild, "not before the next release". Agreed to manually cut heavy replenishment for SYD01 → *no owner assigned*.

AOB: MAD01 complained about late deliveries (lead time 2-3 days vs 1 in the system parameters). Re will check.

Actions
- Marini: PO AMALIA brown small sizes (by 28/09)
- Moreau/Galli: domestic IBT AUREA black
- Conti: reload Golden Week override
- Re: MAD01 lead time
""")

DOCS["2026-10-02_call_ILVARO_pink"] = ("md", """
## Call veloce IRIDIA rosa – ven 02/10 ore 09:15 CET
(Bellini, Chen, Kowalski, Marini, C. Valli marketing/CRM, e-com: T. Rinaldi)

Contesto: ieri sera (01/10) foto di **Yoon Ha-eun (LUMIERE)** fuori dalla sfilata a Parigi con IRIDIA 510. Valli: NON era nel nostro seeding, PR Seoul l'ha scoperto stamattina dai social. Probabilmente acquisto suo o del suo stylist.

- E-com APAC: rosa esaurito in tutte le taglie stamattina tra le 07:20 e le 08:40 ora di Seoul (KR, CN, JP). ~700 richieste "notify me" solo stamattina (Rinaldi, dato provvisorio).
- Negozi Asia: SEL01, SEL02, BUS01 esauriti, TYO in esaurimento. Chen: "it's crazy, waiting lists everywhere".
- DC: zero rosa ovunque fino all'arrivo del top-up giovedì 08/10.
- Discussione split 600 pz (deciso 21/09: HK 330 / IT 190 / US 80). Chen chiede di spostare la quota IT su HK. Marini: la merce arriva dal Vietnam già ripartita per DC, re-routing costa ~1 settimana. **Decisione: split invariato.** Bellini: eventualmente STO IT10→HK30 dopo l'arrivo se l'Europa non vende.
- E-com chiede 150 pz del top-up. Compromesso: riserva web HK 70 / IT 40 / US 15.
- Marini: chiedere subito al fornitore un riordino urgente (capacità limitata, fine novembre?).
- Valli: PR APAC vuole mandare un capo in regalo all'agenzia della cantante → serve 1 pz taglia 0 da HK30.

Prossimo check lunedì all'S&OP.
""")

DOCS["2026-10-05_SOP_FW26_minutes"] = ("docx", """
# S&OP FW26 – weekly meeting
**Date:** Monday 05/10/2026 10:00–12:00 CET
**Attendees:** G. Bellini (chair), P. Conti, L. Marini, M. Re, W. Chen, S. Kowalski, E. Sartori, R. Galli, I. Moreau, F. McLeod, C. Valli, F. Albertini (Trade Compliance, item 6)

## 1. IRIDIA powder pink – after the K-pop effect
- Asia sold out in stores and online between 02/10 and 04/10. Waiting lists SEL01, SEL02, TYO01, SHA01, HKG01, SIN01. Some spill-over on IRIDIA black and off-white.
- Emergency re-order **PO 4500222417** raised today: 1,200 pcs (HK 800 / IT 300 / US 100). Vendor first answer: delivery **27/11**, capacity limited.
- Top-up (600) still landing Thursday 08/10 with split HK 330 / IT 190 / US 80. Web reserve confirmed (HK 70 / IT 40 / US 15).
- PR gifting request for the singer's agency (1 pc, HK30) was rejected by the system (ATP 0). Valli will retry after 08/10.
- Kowalski: the US pink is not moving, happy to release US units to Asia. Bellini: let's see after the top-up, STO lead time 6-7 days.

## 2. IRIDIA black
- DC reserved stock exhausted after the spill-over, black at zero from 06/10. Re-order 4500221188 (850 pcs) ETA 23/10.

## 3. AUREA black
- Vendor moved delivery again (mail 01/10): now **16/10**. First tranche **420 pcs for IT10**, HK30 and US20 lines in parallel (US +3 days). The balance for IT10 (300 pcs) is still on 30/10.
- IBT: Milan got 3 pcs from Rome, Harrods 2 from New Bond St. Paris got nothing (Cannes refused, Geneva cross-border not allowed). Courmayeur has black Aurea in stock but is closed for the season.
- Moreau: "Korean clients will not wait until 16/10, many of them leave after Chuseok and Paris FW".

## 4. Couture ASTRA – week 1-2
- APAC strong (CN/HK ~half of the launch qty sold in 10 days), EMEA weak and US very weak. Feedback from stores: fit runs small, price too high. QN opened (300004602, fit check).
- Proposal Chen: move black ASTRA from IT10 to HK30 (STO). **Decision: approved for ~45% of IT10 free stock, Re to prepare for 15/10.**
- Bellini admits the allocation assumption (EU like the 2025 Couture drop / CN profile) was wrong.

## 5. VEL vest – zipper defect
- Quality notification 300004587 opened 30/09: slider detaches, batch 2607-A114 (delivered from mid-August). ~230 defect returns from stores so far.
- Usage decision 02/10: **batch blocked** in all DCs; good batch 2609-B027 (received 18-21/09) released. Stores to send returns to QM01 (IT10/US20) or QC-HOLD at TKO.
- Open question (Sartori): can stores keep selling 2607 units already on the floor? No answer from Quality yet.

## 6. Logistics
- MAD01 mispick (navy NOVA instead of black, two deliveries): root cause bin labels swapped at CSG; cycle count 06/10. Black NOVA in Thursday's delivery.
- Albertini: Korea import document checks getting stricter for down garments; nothing blocked so far.

## 7. Other
- Japan: warm weather (27-29 °C) – heavy outerwear not moving in TYO/OSA. Price increase letter to clients goes out 09/10.
- VLOGO CARDIGAN size M: DC sold out since 01/10, knit re-order 4500221905 wk44 (30/10). Galli: MIL01 has only 1-2 left.
- SYD01: no news. ROM01 traffic -50% with the scaffolding, Galli asks to lower ROM01 cover → **not done**, parameters frozen until November.

## Action items
| Owner | Action | Due |
|---|---|---|
| Marini | Push vendor for earlier slot on 4500222417 | 09/10 |
| Re | STO ASTRA IT10→HK30 | 15/10 |
| Bellini | Allocation plan for the 420 AUREA black (EMEA) | 12/10 |
| Valli | Retry PR gifting IRIDIA pink after 08/10 | 09/10 |
| Quality (n/a) | Rule for VEL 2607 units on the shop floor | ? |
""")

DOCS["2026-10-19_SOP_FW26_minutes"] = ("docx", """
# S&OP FW26 – riunione settimanale
**Data:** lunedì 19/10/2026 – 10:00 CET
**Presenti:** Bellini, Conti, Marini, Re, Chen, Kowalski, Sartori, Galli, Moreau, McLeod, Lim, Valli

## 1. AUREA nero – arrivato
- Ricevuti venerdì 16/10: IT10 420 pz, HK30 1.040 pz. US20 ricevuto oggi (220). Il saldo IT10 (300 pz) anticipato e arrivato stamattina.
- Allocazione EMEA (proposta Bellini): ~45% su PAR01/PAR02/MIL01/LON02, resto pro-quota. Spedizione domani 20/10.
- Moreau: i clienti coreani sono calati molto da metà mese. Bellini: "manteniamo la priorità, la domanda c'è ancora". *(da verificare: rischio di sovra-allocare Parigi/Milano ora che il picco è passato)*
- HK30: spediti 238 pz il 17/10 a SEL/TYO/SHA ecc., seconda ondata il 21/10.

## 2. IRIDIA rosa
- Top-up spedito ai negozi dal 09/10. SEL01: il primo carico era bloccato alla dogana di Incheon (certificato KC mancante, 08→13/10) e conteneva solo off-white; il rosa (20 pz) è partito il 10/10 ed è stato a sua volta trattenuto per ispezione → consegnato **18/10** invece del 15/10.
- Riordino 4500222417: il fornitore ha confermato il **20/11** (prima 27/11). Niente altro prima.
- In DC resta ancora rosa: circa 110 pz a HK30 (venerdì) e ~80 a IT10. Chen chiede perché non sono stati spediti ai negozi asiatici che hanno le liste d'attesa. Conti: il riordino automatico copre il target di stock "normale" (forecast senza evento), quindi non ha richiesto di più. **Nessuna decisione presa**, rimandato.
- Valli: il regalo all'agenzia della cantante non è ancora partito.

## 3. IRIDIA nero
- Riordino 4500221188 anticipato al 20/10 (aereo). Da domani torna disponibile in tutte le DC.

## 4. Freddo DACH (12-19/10)
- Prima neve a Monaco 14/10, vendite heavy esplose MUN01/ZRH01/BER01/GVA01/AMS01/CPH01. Alcune taglie a zero su MONTELUCE/COSIMO. IBT da BER01 a MUN01 fatto.

## 5. VEL
- 2607 bloccato in DC (~1.100 pz tra le tre DC). Resi difettosi arrivati a QM01/QC-HOLD ~350 pz. 8D del fornitore non ancora ricevuto. Regola per i pezzi 2607 nei negozi: ancora aperta.

## 6. ASTRA
- STO IT10→HK30 partito 15/10, arrivo previsto a TKO 20/10. Cina continua bene, inizia la prevendita Double 11 online da domani (clienti chiedono di allineare i prezzi Tmall in negozio).

## 7. Varie
- VLOGO CARDIGAN M: fornitore conferma 60% il 30/10 e 40% il 06/11. MIL01 a zero da giovedì 08/10, IBT da Rinascente approvato ma non ancora uscito.
- AMALIA marrone 0-2: PO 4500222090 confermato 10/11.
- Giappone: corsa agli acquisti prima dell'aumento del 1/11, icone (NOVA, AUREA) molto forti da 12/10.
- STM01 chiuso da oggi (Zwischensaison) – **bloccare i riordini** (parametri ancora attivi!). CRT01 solo weekend.
- SYD01: Lim di nuovo. Conti propone di passare il profilo stagionale a SH "dal prossimo run" e di tagliare manualmente heavy → owner Conti, nessuna data.

## Action items
| # | Azione | Owner | Scadenza |
|---|--------|-------|----------|
| 1 | Decidere destinazione rosa residuo HK30/IT10 | Bellini / Chen | 21/10 |
| 2 | Bloccare riordini STM01 | Conti | 19/10 |
| 3 | Profilo SH per SYD01 | Conti | ? |
| 4 | Regola VEL 2607 in negozio | Quality | ? |
| 5 | Sollecito 8D VEL | Marini | 23/10 |
""")

# draft version of the Oct 19 minutes, circulated before corrections (version mismatch with the docx)
DOCS["2026-10-19_SOP_FW26_minutes_DRAFT"] = ("md", """
# S&OP FW26 - 19/10 - BOZZA (non rivista)

1) AUREA nero arrivato 16/10. IT10 420, HK 1040, US arriva 19/10. Saldo IT 300 pz ancora al 30/10.
   Allocazione EMEA: 60% PAR01/PAR02/MIL01/LON02. Spedizione 20/10.
2) IRIDIA rosa: riordino 4500222417 al 27/11. In DC: HK ~110, IT ~80. Da decidere.
   SEL01 rosa arrivato 13/10 (dogana).
3) IRIDIA nero riordino 23/10.
4) Freddo DACH ok.
5) VEL bloccato ~1.100 pz.
6) Astra STO arrivato.
7) VLOGO CARDIGAN M 30/10. AMALIA 10/11. STM01 chiuso. SYD01 -> Conti.
""")

DOCS["planner_running_notes_GBellini"] = ("md", """
note mie – FW26 (G.B.)

03/09 lancio ok. amalia marrone gira!! split 70/30 sbagliato
07/09 aurea nero DC fino al 12 circa. PO 25/9 (Nistru). Marini dice tranquillo
10/09 mail Nistru: 2/10. tessuto laccato rifiutato in tintoria. ok 2 settimane di buco
14/09 OTB v5 mandato a Sartori. Astra EU = CN, speriamo
17/09 HKG chiuso per tifone, niente di grave
21/09 split top-up IRIDIA rosa HK330/IT190/US80 - rosa vende poco, forse troppo anche così
22/09 Galli: coreane a Monte Napo per Aurea nera, ROM01 manda 3 pz. Parigi niente.
24/09 lancio Astra. vedremo
25/09 VIC MIL01 vuole Amalia marrone 1 -> DC zero. VEN01 ne ha uno, IBT
30/09 QN VEL aperta da quality, cursore zip
01/10 Nistru: aurea nero 16/10 (!!). 420 su IT. il resto 30/10
02/10 !!! cantante coreana con IRIDIA rosa a Parigi. web asia esaurito in un'ora. split NON cambiato (Marini: rerouting 1 sett)
02/10 VEL 2607 bloccato
05/10 riordino rosa 1200 pz, fornitore 27/11. troppo tardi?
06/10 inventario CSG trova +37 nova nero (etichette scambiate). MAD01 ha 37 blu che non vende
08/10 top-up rosa arrivato. SEL01 dogana Incheon (KC??) - Albertini non sapeva del requisito
08/10 MIL01 cardigan logo M finito. DC zero fino 30/10
09/10 lettera aumento prezzi JP
12/10 freddo DE/CH. bene
14/10 Nistru anticipa il saldo 300 aurea IT al 19/10
16/10 aurea nero arrivato. Astra STO partito ieri
16/10 fornitore rosa conferma 20/11
19/10 S&OP. in DC ancora rosa (HK ~110, IT ~80): perché non è uscito? il riordino auto non lo chiede, target senza evento. da sistemare
20/10 double 11 prevendite, SHA chiede price match. no.
21/10 iridia nero ritornato in DC (20/10). EMEA ordini giovedì
""")


def md_to_docx(md, path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    lines = md.strip().splitlines()
    table_buf = []

    def flush_table():
        if not table_buf:
            return
        rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in table_buf if not re.match(r"^\|\s*-", l)]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                t.cell(i, j).text = v
        table_buf.clear()

    for l in lines:
        if l.startswith("|"):
            table_buf.append(l)
            continue
        flush_table()
        if l.startswith("# "):
            doc.add_heading(l[2:], level=1)
        elif l.startswith("## "):
            doc.add_heading(l[3:], level=2)
        elif l.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, l[2:])
        elif l.strip():
            p = doc.add_paragraph()
            add_runs(p, l)
    flush_table()
    doc.core_properties.author = "Giorgio Bellini"
    doc.save(path)


def add_runs(p, text):
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and len(part) > 1:
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


for name, (fmt, text) in DOCS.items():
    if fmt == "docx":
        md_to_docx(text, OUT / f"{name}.docx")
    else:
        (OUT / f"{name}.md").write_text(text.strip() + "\n", encoding="utf-8")
print("docs:", sorted(p.name for p in OUT.iterdir()))
